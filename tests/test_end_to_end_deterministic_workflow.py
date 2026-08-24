"""End-to-end API coverage for the controlled workflow without a live model.

The deterministic agent below is a test double for the model boundary.  It
only supplies the structured output expected from a successful model call; all
authorisation, persistence, evidence validation, approvals and PowerPoint
export still execute through the real HTTP API and SQLite repository.
"""

from io import BytesIO
import time

import fitz
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage
from pptx import Presentation as PowerPoint

from app.ai.workflows.graph_state import GraphState
from app.domain.enums.workflow_status import WorkflowStatus
from app.domain.enums.workflow_step import WorkflowStep
from app.domain.models.agenda import Agenda
from app.domain.models.blueprint import Blueprint
from app.domain.models.execution_context import ExecutionContext
from app.domain.models.resource_analysis import ResourceAnalysis
from app.domain.models.slide import Slide
from app.domain.models.slide_outline import SlideOutline
from app.interfaces.api import main as api
from app.interfaces.api.routers import jobs as job_routes
from app.interfaces.storage.user_session_repository import UserSessionRepository


class DeterministicPresentationAgent:
    """A local, no-network substitute for two successful model turns."""

    def __init__(self) -> None:
        self.calls = 0

    def invoke(self, state: GraphState, thread_id: str) -> dict:
        del thread_id
        self.calls += 1
        result = state.model_copy(deep=True)
        presentation = result.presentation
        assert presentation is not None
        resource = presentation.resources[0]
        evidence_excerpt = "Vitamin D status should be assessed using the supplied clinical evidence."

        if presentation.blueprint is None:
            presentation.blueprint = Blueprint(
                title=presentation.title,
                learning_objective=presentation.context.objective,
                target_number_of_slides=2,
                storytelling="Evidence first, then clinical application.",
                sections=["Evidence", "Application"],
                slides=[
                    SlideOutline(
                        slide_number=1,
                        title="Evidence overview",
                        objective="Introduce the supplied evidence.",
                        key_message="Use the uploaded evidence as the sole source.",
                    ),
                    SlideOutline(
                        slide_number=2,
                        title="Clinical application",
                        objective="Discuss the evidence in the intended context.",
                        key_message="Keep claims within the uploaded resource.",
                    ),
                ],
            )
            presentation.agenda = Agenda(items=["Evidence overview", "Clinical application"])
            presentation.state.current_step = WorkflowStep.BLUEPRINT_VALIDATION
            presentation.state.workflow_status = WorkflowStatus.AWAITING_AGENDA_APPROVAL
            result.execution = ExecutionContext(last_tool="build_blueprint")
            result.messages.append(AIMessage(content="Deterministic blueprint ready for human review."))
            return result.model_dump()

        presentation.slides = [
            Slide(
                slide_number=outline.slide_number,
                title=outline.title,
                objective=outline.objective,
                key_messages=[outline.key_message],
                content=outline.key_message,
                references=[f"{resource.title}, p. 1"],
                reference_details=[
                    {
                        "resource_id": resource.id,
                        "page": 1,
                        "evidence_excerpt": evidence_excerpt,
                        "reference": resource.title,
                    }
                ],
            )
            for outline in presentation.blueprint.slides
        ]
        presentation.state.current_step = WorkflowStep.SLIDE_VALIDATION
        presentation.state.total_slides = len(presentation.slides)
        presentation.state.current_slide = len(presentation.slides)
        presentation.state.workflow_status = WorkflowStatus.AWAITING_SLIDE_APPROVAL
        result.execution = ExecutionContext(last_tool="generate_slides")
        result.messages.append(AIMessage(content="Deterministic slides ready for human review."))
        return result.model_dump()


class DeterministicBlueprintGeneration:
    """Model boundary substitute used only by the explicit blueprint command."""

    def __init__(self, agent: DeterministicPresentationAgent) -> None:
        self.agent = agent

    def execute(self, state: GraphState) -> GraphState:
        return GraphState(**self.agent.invoke(state, "deterministic-blueprint"))


class DeterministicSlideGeneration:
    """Model boundary substitute used only by the explicit slide command."""

    def __init__(self, agent: DeterministicPresentationAgent) -> None:
        self.agent = agent

    def execute(self, state: GraphState) -> GraphState:
        return GraphState(**self.agent.invoke(state, "deterministic-slides"))


class DeterministicSlideRegeneration:
    """No-network substitute that proves the persisted reviewer feedback is used."""

    def __init__(self) -> None:
        self.received_comments: list[str | None] = []

    def execute(self, state: GraphState, index: int) -> GraphState:
        presentation = state.presentation
        assert presentation is not None
        slide = presentation.slides[index]
        self.received_comments.append(slide.reviewer_comments)
        slide.content = "Regenerated from the HCP reviewer request."
        slide.is_validated = False
        return state


class DeterministicBlueprintRegeneration:
    """No-network substitute proving all typed blueprint feedback is durable."""

    def __init__(self) -> None:
        self.received_comments: list[str | None] = []

    def execute(self, state: GraphState) -> GraphState:
        presentation = state.presentation
        assert presentation is not None and presentation.blueprint is not None
        self.received_comments = [outline.reviewer_comments for outline in presentation.blueprint.slides]
        presentation.blueprint.is_validated = False
        presentation.state.blueprint_validated = False
        return state


class DeterministicResourceOverview:
    """Local substitute for the Resource Overview model boundary."""

    def execute(self, resources, **kwargs) -> ResourceAnalysis:
        del kwargs
        resource = resources[0]
        return ResourceAnalysis(
            summary=(
                "Overall idea: the uploaded PDF supports a focused Vitamin D discussion "
                f"[{resource.id}, p. 1]."
            ),
            resource_ids=[resource.id],
        )


class DeterministicResourceDiscussion:
    """Local substitute for the Resource Chat model boundary."""

    def execute(self, resources, question, **kwargs) -> str:
        del question, kwargs
        resource = resources[0]
        return f"The supplied PDF supports the requested discussion [{resource.id}, p. 1]."


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Vitamin D status should be assessed using the supplied clinical evidence. "
        "This source is provided by the user for the presentation workflow.",
    )
    content = document.tobytes()
    document.close()
    return content


def _wait_for_job(client: TestClient, headers: dict[str, str], user_id: str, job_id: str) -> dict:
    """Poll the real local job API without calling an external provider."""
    for _ in range(100):
        response = client.get(f"/api/v1/jobs/{user_id}/{job_id}", headers=headers)
        assert response.status_code == 200
        job = response.json()
        if job["status"] in {"completed", "failed"}:
            return job
        time.sleep(0.01)
    raise AssertionError("The local workflow job did not reach a terminal state.")


def test_end_to_end_human_controlled_workflow_without_live_model(tmp_path, monkeypatch):
    """Cover the HCP path from account creation to a valid PPTX download."""
    repository = UserSessionRepository(tmp_path / "hpa-e2e.sqlite3")
    deterministic_agent = DeterministicPresentationAgent()
    monkeypatch.setattr(api, "get_repository", lambda: repository)
    monkeypatch.setattr(api, "SummarizeResourcesUseCase", DeterministicResourceOverview)
    monkeypatch.setattr(api, "DiscussResourcesUseCase", DeterministicResourceDiscussion)
    monkeypatch.setattr(
        job_routes,
        "BuildBlueprintWorkflowUseCase",
        lambda: DeterministicBlueprintGeneration(deterministic_agent),
    )
    monkeypatch.setattr(
        job_routes,
        "GenerateSlidesWorkflowUseCase",
        lambda: DeterministicSlideGeneration(deterministic_agent),
    )
    client = TestClient(api.app)

    registered = client.post(
        "/auth/register",
        json={
            "user_id": "hcp-e2e",
            "password": "safe-local-password",
            "professional_role": "specialist_physician",
            "preferred_language": "en",
        },
    )
    assert registered.status_code == 200
    headers = {"Authorization": f"Bearer {registered.json()['token']}"}

    assert client.post(
        "/projects",
        headers=headers,
        json={"user_id": "hcp-e2e", "project_id": "vitamin-d", "project_name": "Vitamin D update"},
    ).status_code == 200

    setup = client.post(
        "/projects/hcp-e2e/vitamin-d/presentation/setup",
        headers=headers,
        json={
            "topic": "Vitamin D clinical update",
            "audience": "general_practitioner",
            "presentation_type": "Lecture",
            "language": "English",
            "duration_minutes": 10,
            "objective": "Review the supplied clinical evidence.",
            "presenter_name": "Dr Ada Martin",
            "presenter_title": "Rheumatologist",
            "organization": "Teaching Hospital",
            "event_name": "HCP Update 2026",
            "venue": "Algiers",
            "presentation_date": "9 August 2026",
        },
    )
    assert setup.status_code == 200
    assert setup.json()["presentation"]["context"]["presenter_name"] == "Dr Ada Martin"

    uploaded = client.post(
        "/resources/pdf/hcp-e2e/vitamin-d",
        headers=headers,
        files={"file": ("vitamin-d-evidence.pdf", _pdf_bytes(), "application/pdf")},
    )
    assert uploaded.status_code == 200
    resource_id = uploaded.json()["resource_id"]

    # Resource Overview and Resource Chat are intentionally independent from
    # production. They can be used immediately after upload and must keep
    # their own transcript without creating a presentation-chat message.
    overview_job = client.post(
        "/api/v1/resources/hcp-e2e/vitamin-d/overview/jobs", headers=headers
    )
    assert overview_job.status_code == 202
    assert _wait_for_job(client, headers, "hcp-e2e", overview_job.json()["job_id"])["status"] == "completed"
    overview_project = client.get("/projects/hcp-e2e/vitamin-d", headers=headers).json()
    assert "Vitamin D discussion" in overview_project["resource_analysis"]["summary"]
    assert overview_project["messages"] == []

    resource_chat_job = client.post(
        "/api/v1/resources/hcp-e2e/vitamin-d/discussion/jobs",
        headers=headers,
        json={"question": "What does this PDF support?"},
    )
    assert resource_chat_job.status_code == 202
    assert _wait_for_job(client, headers, "hcp-e2e", resource_chat_job.json()["job_id"])["status"] == "completed"
    resource_chat_project = client.get("/projects/hcp-e2e/vitamin-d", headers=headers).json()
    assert [turn["role"] for turn in resource_chat_project["resource_messages"]] == ["user", "assistant"]
    assert resource_chat_project["messages"] == []

    assert client.post(
        f"/projects/hcp-e2e/vitamin-d/resources/{resource_id}/attach", headers=headers
    ).status_code == 200
    validated_resources = client.post(
        "/projects/hcp-e2e/vitamin-d/resources/validate", headers=headers
    )
    assert validated_resources.status_code == 200
    assert validated_resources.json()["presentation"]["state"]["workflow_status"] == "blueprint_generation"

    blueprint_job = client.post(
        "/api/v1/projects/hcp-e2e/vitamin-d/blueprint/jobs",
        headers=headers,
    )
    assert blueprint_job.status_code == 202
    assert _wait_for_job(client, headers, "hcp-e2e", blueprint_job.json()["job_id"])["status"] == "completed"

    deterministic_blueprint_regeneration = DeterministicBlueprintRegeneration()
    monkeypatch.setattr(
        job_routes,
        "RegenerateBlueprintUseCase",
        lambda: deterministic_blueprint_regeneration,
    )
    regenerated_blueprint = client.post(
        "/api/v1/projects/hcp-e2e/vitamin-d/blueprint/regenerate/jobs",
        headers=headers,
        json={"comments_by_index": {"0": "Make the evidence opening more concise for this audience."}},
    )
    assert regenerated_blueprint.status_code == 202
    assert _wait_for_job(client, headers, "hcp-e2e", regenerated_blueprint.json()["job_id"])["status"] == "completed"
    assert deterministic_blueprint_regeneration.received_comments == [
        "Make the evidence opening more concise for this audience.",
        None,
    ]

    assert client.post("/projects/hcp-e2e/vitamin-d/agenda/approve", headers=headers).status_code == 200
    for index in (0, 1):
        assert client.post(
            f"/projects/hcp-e2e/vitamin-d/blueprint/items/{index}/approve",
            headers=headers,
            json={"comments": "Reviewed by the HCP."},
        ).status_code == 200
    assert client.post("/projects/hcp-e2e/vitamin-d/blueprint/approve", headers=headers).status_code == 200

    slides_job = client.post(
        "/api/v1/projects/hcp-e2e/vitamin-d/slides/jobs",
        headers=headers,
    )
    assert slides_job.status_code == 202
    assert _wait_for_job(client, headers, "hcp-e2e", slides_job.json()["job_id"])["status"] == "completed"

    # Regeneration is also a durable job. The typed reviewer comment must be
    # persisted before the worker runs and made available to the use case.
    deterministic_regeneration = DeterministicSlideRegeneration()
    monkeypatch.setattr(
        job_routes,
        "RegenerateSlideUseCase",
        lambda: deterministic_regeneration,
    )
    regenerated = client.post(
        "/api/v1/projects/hcp-e2e/vitamin-d/slides/0/regenerate/jobs",
        headers=headers,
        json={"comments": "Make the first slide more concise for the audience."},
    )
    assert regenerated.status_code == 202
    assert _wait_for_job(client, headers, "hcp-e2e", regenerated.json()["job_id"])["status"] == "completed"
    assert deterministic_regeneration.received_comments == [
        "Make the first slide more concise for the audience."
    ]
    regenerated_project = client.get("/projects/hcp-e2e/vitamin-d", headers=headers).json()
    assert regenerated_project["presentation"]["slides"][0]["content"] == (
        "Regenerated from the HCP reviewer request."
    )

    # A direct HCP rewrite is an API-only operation: no extra model call and
    # no inherited AI-evidence label for the new human-authored text.
    edited = client.put(
        "/projects/hcp-e2e/vitamin-d/slides/1",
        headers=headers,
        json={
            "title": "HCP concluding message",
            "objective": "Close the presentation.",
            "key_messages": ["This conclusion was written by the HCP."],
            "content": "This conclusion was written by the HCP.",
            "speaker_notes": "Close and invite discussion.",
            "content_origin": "user_authored",
        },
    )
    assert edited.status_code == 200
    assert edited.json()["presentation"]["slides"][1]["evidence_verified"] is False
    assert deterministic_agent.calls == 2


    for index in (0, 1):
        assert client.post(
            f"/projects/hcp-e2e/vitamin-d/slides/{index}/approve",
            headers=headers,
            json={"comments": "Approved by the HCP."},
        ).status_code == 200
    assert client.post("/projects/hcp-e2e/vitamin-d/slides/approve", headers=headers).status_code == 200
    assert client.post("/projects/hcp-e2e/vitamin-d/presentation/approve", headers=headers).status_code == 200

    exported = client.get("/presentations/hcp-e2e/vitamin-d/export/pptx", headers=headers)
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    deck = PowerPoint(BytesIO(exported.content))
    title_slide_text = " ".join(shape.text for shape in deck.slides[0].shapes if hasattr(shape, "text"))
    assert len(deck.slides) == 5  # title, agenda, two slides, mandatory resources slide
    assert "Dr Ada Martin" in title_slide_text
    assert "HCP Update 2026" in title_slide_text
    assert deterministic_agent.calls == 2


def test_scope_declaration_api_requires_explicit_human_fields(tmp_path, monkeypatch):
    """The chat model cannot resolve a scope mismatch on the user's behalf."""
    repository = UserSessionRepository(tmp_path / "hpa-scope.sqlite3")
    monkeypatch.setattr(api, "get_repository", lambda: repository)
    client = TestClient(api.app)

    registered = client.post(
        "/auth/register",
        json={
            "user_id": "scope-hcp",
            "password": "safe-local-password",
            "professional_role": "veterinarian",
            "preferred_language": "en",
        },
    )
    headers = {"Authorization": f"Bearer {registered.json()['token']}"}
    assert client.post(
        "/projects",
        headers=headers,
        json={"user_id": "scope-hcp", "project_id": "scope-project"},
    ).status_code == 200
    assert client.post(
        "/projects/scope-hcp/scope-project/presentation/setup",
        headers=headers,
        json={
            "topic": "Human medicine evidence update",
            "audience": "general_practitioner",
            "presentation_type": "Lecture",
            "language": "English",
            "duration_minutes": 10,
            "objective": "Review user-provided evidence for healthcare professionals.",
        },
    ).status_code == 200
    thread_id, state = repository.load("scope-hcp", "scope-project")
    assert state.presentation is not None
    state.presentation.state.workflow_status = WorkflowStatus.AWAITING_SCOPE_CLARIFICATION
    api._save_project("scope-hcp", "scope-project", thread_id, state, event_type="TEST_SCOPE_BLOCK")

    rejected = client.post(
        "/projects/scope-hcp/scope-project/presentation/scope-clarification",
        headers=headers,
        json={
            "declared_role": "Medical representative",
            "delivery_purpose": "I present evidence to healthcare professionals.",
            "confirmed_within_scope": False,
        },
    )
    assert rejected.status_code == 409
    assert rejected.json()["detail"]["code"] == "SCOPE_CONFIRMATION_REQUIRED"

    accepted = client.post(
        "/projects/scope-hcp/scope-project/presentation/scope-clarification",
        headers=headers,
        json={
            "declared_role": "Medical representative with veterinary training",
            "delivery_purpose": "I present human-health scientific information to healthcare professionals.",
            "confirmed_within_scope": True,
        },
    )
    assert accepted.status_code == 200
    declaration = accepted.json()["presentation"]["professional_scope_declaration"]
    assert declaration["declared_role"] == "Medical representative with veterinary training"
    assert accepted.json()["presentation"]["state"]["workflow_status"] == "blueprint_generation"
