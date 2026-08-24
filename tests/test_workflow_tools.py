import pytest

from app.ai.workflows.graph_state import GraphState
from app.ai.workflows.tools import (
    collect_context,
    create_presentation,
    validate_context,
    validate_final_presentation,
    validate_slides,
)
from app.domain.enums.audience_type import AudienceType
from app.domain.enums.language import Language
from app.domain.enums.presentation_type import PresentationType
from app.domain.models.slide import Slide
from app.domain.models.slide_outline import SlideOutline
from app.domain.models.blueprint import Blueprint
from app.application.use_cases.workflow_steps import (
    AuthorSlideFromBlueprintUseCase,
    EditBlueprintItemUseCase,
    EditSlideUseCase,
    ReviewSlideUseCase,
)
from app.application.use_cases.workflow_steps import (
    BuildBlueprintWorkflowUseCase,
    RecordProfessionalScopeUseCase,
)
from app.application.use_cases.export_powerpoint import ExportPowerPointUseCase
from app.application.use_cases.remove_resource import RemoveResourceUseCase
from app.application.use_cases.update_presentation_details import UpdatePresentationDetailsUseCase
from app.application.use_cases.generate_slides import GenerateSlidesUseCase
from app.application.use_cases.add_resource import AddResourceUseCase
from app.application.services.workflow_policy import WorkflowPolicy
from app.ai.agents.healthcare_presentation_agent import HealthcarePresentationAgent
from app.domain.enums.resource_type import ResourceType
from app.domain.models.resource import Resource
from app.domain.models.agenda import Agenda
from app.domain.models.user_profile import UserProfile
from app.domain.enums.workflow_status import WorkflowStatus
from app.domain.exceptions.workflow_error import WorkflowError
from pptx import Presentation as PowerPoint


def test_context_can_create_a_presentation_without_a_live_llm():
    state = collect_context.func(
        GraphState(user_profile=UserProfile(professional_role="biologist", preferred_language="ar")),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review novel treatment strategies",
    )

    state = validate_context.func(state)
    state = create_presentation.func(state)

    assert state.presentation is not None
    assert state.presentation.title == "Depression"
    assert state.presentation.owner_profile.professional_role == "biologist"


def test_optional_title_slide_details_are_preserved_without_blocking_context_validation():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=20,
        objective="Review treatment strategies",
        presenter_name="Dr Ada Martin",
        presenter_title="Associate Professor of Psychiatry",
        organization="University Hospital",
        event_name="Clinical Update 2026",
        venue="Algiers",
        presentation_date="29 July 2026",
    )

    presentation = create_presentation.func(validate_context.func(state)).presentation

    assert presentation.context.presenter_name == "Dr Ada Martin"
    assert presentation.context.event_name == "Clinical Update 2026"
    assert presentation.context.presentation_date == "29 July 2026"


def test_updating_title_slide_details_reopens_only_final_approval():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=20,
        objective="Review treatment strategies",
    )
    state = create_presentation.func(validate_context.func(state))
    state.presentation.state.workflow_status = WorkflowStatus.READY_FOR_EXPORT
    state.presentation.state.presentation_validated = True

    UpdatePresentationDetailsUseCase().execute(
        state,
        presenter_name="Dr Ada Martin",
        event_name="Clinical Update 2026",
    )

    assert state.presentation.context.presenter_name == "Dr Ada Martin"
    assert state.presentation.state.workflow_status == WorkflowStatus.AWAITING_FINAL_APPROVAL
    assert not state.presentation.state.presentation_validated


def test_final_approval_requires_slide_approval_first():
    state = GraphState()
    state.presentation = type("Presentation", (), {"slides": [Slide(slide_number=1, title="Test")], "state": type("State", (), {"slides_validated": False, "presentation_validated": False, "workflow_status": WorkflowStatus.AWAITING_SLIDE_APPROVAL})()})()

    state = ReviewSlideUseCase().execute(state, index=0, comments="Approved")
    state = validate_slides.func(state, approved=True)
    state = validate_final_presentation.func(state, approved=True)

    assert state.presentation.state.slides_validated
    assert state.presentation.state.presentation_validated


def test_veterinarian_profile_requires_scope_clarification_before_blueprint_generation():
    state = collect_context.func(
        GraphState(user_profile=UserProfile(professional_role="veterinarian", preferred_language="fr")),
        topic="Depression",
        audience=AudienceType.GENERAL_PRACTITIONER,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review treatment guidelines",
    )
    state = create_presentation.func(validate_context.func(state))
    state.presentation.resources = [
        Resource(
            id="apa-pdf",
            filename="apa-depression-guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "Depression treatment guideline for primary care."}],
            is_validated=True,
        )
    ]
    state.presentation.state.resources_validated = True
    state.presentation.state.workflow_status = WorkflowStatus.BLUEPRINT_GENERATION

    with pytest.raises(WorkflowError, match="human healthcare audience"):
        BuildBlueprintWorkflowUseCase().execute(state)

    assert state.presentation.state.workflow_status == WorkflowStatus.AWAITING_SCOPE_CLARIFICATION
    state = RecordProfessionalScopeUseCase().execute(
        state,
        declared_role="Medical representative with veterinary training",
        delivery_purpose="I present human-health scientific information to healthcare professionals.",
        confirmed_within_scope=True,
    )

    assert state.presentation.professional_scope is not None
    assert state.presentation.professional_scope_declaration is not None
    assert state.presentation.state.workflow_status == WorkflowStatus.BLUEPRINT_GENERATION


def test_scope_clarification_resumes_slide_generation_when_blueprint_was_already_approved():
    state = GraphState()
    state.presentation = type(
        "Presentation",
        (),
        {
            "state": type(
                "State",
                (),
                {
                    "workflow_status": WorkflowStatus.AWAITING_SCOPE_CLARIFICATION,
                    "blueprint_validated": True,
                },
            )(),
            "professional_scope": None,
        },
    )()

    state = RecordProfessionalScopeUseCase().execute(
        state,
        declared_role="Medical representative with veterinary training",
        delivery_purpose="I present this human-health topic to healthcare professionals.",
        confirmed_within_scope=True,
    )

    assert state.presentation.state.workflow_status == WorkflowStatus.SLIDE_GENERATION


def test_scope_clarification_is_not_accepted_from_chat_text():
    """Only the explicit scope form can resume a blocked workflow."""
    state = collect_context.func(
        GraphState(user_profile=UserProfile(professional_role="veterinarian", preferred_language="en")),
        topic="Depression",
        audience=AudienceType.GENERAL_PRACTITIONER,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=20,
        objective="Review approved product evidence",
    )
    state = create_presentation.func(validate_context.func(state))
    state.presentation.state.workflow_status = WorkflowStatus.AWAITING_SCOPE_CLARIFICATION
    response = HealthcarePresentationAgent._scope_clarification_message_if_needed(state)

    assert response is not None
    assert state.presentation.professional_scope is None
    assert state.presentation.state.workflow_status == WorkflowStatus.AWAITING_SCOPE_CLARIFICATION


def test_scope_clarification_prompt_is_deterministic_and_concise():
    state = GraphState(user_profile=UserProfile(professional_role="veterinarian", preferred_language="en"))
    state.presentation = type(
        "Presentation",
        (),
        {"state": type("State", (), {"workflow_status": WorkflowStatus.AWAITING_SCOPE_CLARIFICATION})()},
    )()

    prompt = HealthcarePresentationAgent._scope_clarification_message_if_needed(state)

    assert prompt is not None
    assert "Professional scope form" in prompt
    assert "clinical guidance for human patients" not in prompt


def test_scope_declaration_requires_human_confirmation():
    state = GraphState()
    state.presentation = type(
        "Presentation",
        (),
        {
            "state": type(
                "State",
                (), {
                    "workflow_status": WorkflowStatus.AWAITING_SCOPE_CLARIFICATION,
                    "blueprint_validated": False,
                },
            )(),
            "professional_scope": None,
        },
    )()

    with pytest.raises(WorkflowError, match="Confirm that this presentation"):
        RecordProfessionalScopeUseCase().execute(
            state,
            declared_role="Medical representative",
            delivery_purpose="I present evidence to healthcare professionals.",
            confirmed_within_scope=False,
        )


def test_deleting_a_resource_resets_generated_content_and_approvals():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review treatment guidelines",
    )
    presentation = create_presentation.func(validate_context.func(state)).presentation
    presentation.resources = [
        Resource(
            id="resource-1",
            filename="guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "Guideline evidence."}],
            is_validated=True,
        )
    ]
    presentation.state.resources_validated = True
    presentation.state.blueprint_validated = True
    presentation.state.slides_validated = True
    presentation.state.presentation_validated = True
    presentation.slides = [Slide(slide_number=1, title="Evidence")]

    RemoveResourceUseCase().execute(presentation, "resource-1")

    assert presentation.resources == []
    assert presentation.blueprint is None
    assert presentation.slides == []
    assert not presentation.state.resources_validated
    assert presentation.state.workflow_status == WorkflowStatus.AWAITING_RESOURCE_UPLOAD


def test_adding_a_resource_is_allowed_late_and_resets_dependent_output():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review treatment guidelines",
    )
    presentation = create_presentation.func(validate_context.func(state)).presentation
    presentation.blueprint = Blueprint(
        title="Blueprint",
        learning_objective="Objective",
        target_number_of_slides=1,
        storytelling="Story",
        slides=[SlideOutline(slide_number=1, title="Evidence", objective="Review", key_message="Message")],
    )
    presentation.agenda = Agenda(items=["Evidence"], is_validated=True)
    presentation.slides = [Slide(slide_number=1, title="Generated slide")]
    presentation.state.resources_validated = True
    presentation.state.blueprint_validated = True
    presentation.state.slides_validated = True
    presentation.state.presentation_validated = True
    presentation.state.workflow_status = WorkflowStatus.READY_FOR_EXPORT

    AddResourceUseCase().execute(
        presentation,
        Resource(
            id="resource-new",
            filename="new-guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "New evidence."}],
            is_validated=True,
        ),
    )

    assert [resource.id for resource in presentation.resources] == ["resource-new"]
    assert presentation.blueprint is None
    assert presentation.agenda is None
    assert presentation.slides == []
    assert not presentation.state.resources_validated
    assert presentation.state.workflow_status == WorkflowStatus.AWAITING_RESOURCE_VALIDATION
    assert WorkflowPolicy.requires_resource_validation(presentation)


def test_blueprint_request_surfaces_human_resource_validation_only_when_needed():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review treatment guidelines",
    )
    state = create_presentation.func(validate_context.func(state))
    state.presentation.resources = [
        Resource(
            id="resource-1",
            filename="guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "Guideline evidence."}],
            is_validated=True,
        )
    ]
    state.presentation.state.workflow_status = WorkflowStatus.AWAITING_RESOURCE_VALIDATION

    assert HealthcarePresentationAgent._needs_resource_validation_for_blueprint(
        state, "Génère le blueprint de ma présentation"
    )
    assert not HealthcarePresentationAgent._needs_resource_validation_for_blueprint(
        state, "Résume les points clés de cette ressource"
    )

    state.presentation.state.resources_validated = True
    assert not HealthcarePresentationAgent._needs_resource_validation_for_blueprint(
        state, "Génère le blueprint de ma présentation"
    )


def test_user_blueprint_edit_is_traced_and_invalidates_dependent_workflow_steps():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review treatment guidelines",
    )
    state = create_presentation.func(validate_context.func(state))
    presentation = state.presentation
    presentation.blueprint = Blueprint(
        title="Blueprint",
        learning_objective="Objective",
        target_number_of_slides=1,
        storytelling="Story",
        slides=[SlideOutline(slide_number=1, title="AI title", objective="AI objective", key_message="AI message", is_validated=True)],
        is_validated=True,
    )
    presentation.agenda = Agenda(items=["AI title"], is_validated=True)
    presentation.slides = [Slide(slide_number=1, title="Old generated slide")]
    presentation.state.blueprint_validated = True
    presentation.state.slides_validated = True
    presentation.state.presentation_validated = True
    presentation.state.workflow_status = WorkflowStatus.AWAITING_BLUEPRINT_APPROVAL

    EditBlueprintItemUseCase().execute(
        state,
        0,
        title="User title",
        objective="User objective",
        key_message="User message",
        content_origin="user_authored",
    )

    item = presentation.blueprint.slides[0]
    assert item.content_origin == "user_authored"
    assert item.original_ai_snapshot["title"] == "AI title"
    assert not item.is_validated
    assert not presentation.agenda.is_validated
    assert presentation.agenda.items == ["User title"]
    assert presentation.slides == []
    assert presentation.state.workflow_status == WorkflowStatus.AWAITING_AGENDA_APPROVAL


def test_user_slide_edit_requires_new_human_approval_without_inheriting_ai_evidence():
    state = GraphState()
    state.presentation = type(
        "Presentation",
        (),
        {
            "slides": [
                Slide(
                    slide_number=1,
                    title="AI title",
                    key_messages=["AI message"],
                    content="AI content",
                    evidence_verified=True,
                    is_validated=True,
                )
            ],
            "state": type(
                "State",
                (),
                {
                    "slides_validated": True,
                    "presentation_validated": True,
                    "workflow_status": WorkflowStatus.AWAITING_SLIDE_APPROVAL,
                },
            )(),
        },
    )()

    EditSlideUseCase().execute(
        state,
        0,
        title="Human title",
        objective="",
        key_messages=["Human message"],
        content="Human content",
        speaker_notes="",
        content_origin="user_edited",
    )

    slide = state.presentation.slides[0]
    assert slide.content_origin == "user_edited"
    assert slide.original_ai_snapshot["content"] == "AI content"
    assert not slide.evidence_verified
    assert slide.evidence_review_required
    assert not slide.is_validated
    assert not state.presentation.state.slides_validated


def test_powerpoint_always_ends_with_user_validated_resources(tmp_path):
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.FRENCH,
        duration_minutes=20,
        objective="Review novel treatment strategies",
    )
    state = create_presentation.func(validate_context.func(state))
    presentation = state.presentation
    presentation.slides = [
        Slide(
            slide_number=1,
            title="Introduction",
            key_messages=["Message"],
            reference_details=[
                {
                    "title": "Clinical guideline",
                    "resource_id": "resource-1",
                    "page": 1,
                    "evidence_excerpt": "Evidence supporting clinical management.",
                }
            ],
        )
    ]
    presentation.resources = [
        Resource(
            id="resource-1",
            filename="guideline.pdf",
            title="Clinical guideline",
            source="WHO",
            file_type=ResourceType.PDF,
            extracted_text="Evidence",
            extracted_pages=[{"page": 1, "text": "Evidence supporting clinical management."}],
            is_validated=True,
        )
    ]
    presentation.state.resources_validated = True
    presentation.state.presentation_validated = True
    presentation.agenda = Agenda(items=["Introduction", "Treatment"], is_validated=True)
    presentation.context = presentation.context.model_copy(
        update={
            "presenter_name": "Dr Ada Martin",
            "presenter_title": "Professor of Psychiatry",
            "organization": "University Hospital",
            "event_name": "Clinical Update 2026",
            "venue": "Algiers",
            "presentation_date": "29 July 2026",
        }
    )

    custom_template = tmp_path / "custom-template.pptx"
    base_deck = PowerPoint()
    base_deck.slides.add_slide(base_deck.slide_layouts[0])
    base_deck.save(custom_template)

    path = ExportPowerPointUseCase().execute(presentation, tmp_path, custom_template)
    deck = PowerPoint(path)
    last_slide_text = " ".join(shape.text for shape in deck.slides[-1].shapes if hasattr(shape, "text"))

    assert len(deck.slides) == 4  # title, mandatory agenda, content, resources
    assert deck.slides[1].shapes[1].text == "Agenda"
    assert deck.slides[-1].shapes[1].text == "Ressources et validation"
    assert "guideline" in last_slide_text
    assert "validées par l’utilisateur" in last_slide_text
    title_slide_text = " ".join(shape.text for shape in deck.slides[0].shapes if hasattr(shape, "text"))
    resources_slide_text = " ".join(shape.text for shape in deck.slides[-1].shapes if hasattr(shape, "text"))
    assert "Dr Ada Martin" in title_slide_text
    assert "University Hospital" in title_slide_text
    assert "Clinical Update 2026" in title_slide_text
    assert "Clinical guideline" in resources_slide_text
    assert "Identifiant d’audit: resource-1" in resources_slide_text


def test_powerpoint_splits_many_resource_blocks_over_multiple_final_slides(tmp_path):
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=20,
        objective="Review treatment strategies",
    )
    presentation = create_presentation.func(validate_context.func(state)).presentation
    presentation.slides = [
        Slide(
            slide_number=1,
            title="Introduction",
            key_messages=["Evidence-based message"],
            reference_details=[
                {
                    "title": "Guideline 1",
                    "resource_id": "resource-1",
                    "page": 1,
                    "evidence_excerpt": "Evidence one supports this presentation.",
                }
            ],
        )
    ]
    presentation.resources = [
        Resource(
            id=f"resource-{index}",
            filename=f"guideline-{index}.pdf",
            title=f"A deliberately readable title for guideline {index}",
            source="Professional organisation",
            file_type=ResourceType.PDF,
            extracted_pages=[
                {
                    "page": 1,
                    "text": "Evidence one supports this presentation." if index == 1 else f"Evidence for guideline {index}.",
                }
            ],
            is_validated=True,
        )
        for index in range(1, 6)
    ]
    presentation.agenda = Agenda(items=["Introduction"], is_validated=True)
    presentation.state.resources_validated = True
    presentation.state.presentation_validated = True

    path = ExportPowerPointUseCase().execute(presentation, tmp_path)
    deck = PowerPoint(path)
    resource_slides = [deck.slides[len(deck.slides) - 2], deck.slides[len(deck.slides) - 1]]

    assert len(deck.slides) == 5  # title, agenda, content, two resource slides
    assert "Resources and validation (1/2)" in " ".join(shape.text for shape in resource_slides[0].shapes if hasattr(shape, "text"))
    assert "Audit ID: resource-5" in " ".join(shape.text for shape in resource_slides[1].shapes if hasattr(shape, "text"))


def test_powerpoint_labels_user_authored_slide_without_claiming_verified_evidence(tmp_path):
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=10,
        objective="Review treatment options",
    )
    presentation = create_presentation.func(validate_context.func(state)).presentation
    presentation.slides = [
        Slide(
            slide_number=1,
            title="Clinical reflection",
            key_messages=["Human-authored message"],
            content_origin="user_authored",
            is_validated=True,
        )
    ]
    presentation.resources = [
        Resource(
            id="resource-1",
            filename="guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "Evidence."}],
            is_validated=True,
        )
    ]
    presentation.agenda = Agenda(items=["Clinical reflection"], is_validated=True)
    presentation.state.resources_validated = True
    presentation.state.presentation_validated = True

    path = ExportPowerPointUseCase().execute(presentation, tmp_path)
    deck = PowerPoint(path)
    content_text = " ".join(shape.text for shape in deck.slides[2].shapes if hasattr(shape, "text"))
    resources_text = " ".join(shape.text for shape in deck.slides[-1].shapes if hasattr(shape, "text"))

    assert "User-authored content" in content_text
    assert "Human-approved" in content_text
    assert "1 user-authored" in resources_text


def test_user_authored_blueprint_item_becomes_a_slide_without_a_model_call():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=10,
        objective="Review a user-authored message",
    )
    presentation = create_presentation.func(validate_context.func(state)).presentation
    presentation.blueprint = Blueprint(
        title="Blueprint",
        learning_objective="Objective",
        target_number_of_slides=1,
        storytelling="Story",
        slides=[
            SlideOutline(
                slide_number=1,
                title="My own slide",
                objective="My own objective",
                key_message="Exact user-provided message",
                content_origin="user_authored",
            )
        ],
    )
    presentation.resources = [
        Resource(
            id="resource-1",
            filename="guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "Unrelated source text."}],
            is_validated=True,
        )
    ]

    # No constructor: this branch must not initialize or call an LLM chain.
    result = GenerateSlidesUseCase.execute(object.__new__(GenerateSlidesUseCase), presentation)

    assert result.state.workflow_status == WorkflowStatus.AWAITING_SLIDE_APPROVAL
    assert result.slides[0].content_origin == "user_authored"
    assert result.slides[0].title == "My own slide"
    assert result.slides[0].key_messages == ["Exact user-provided message"]
    assert not result.slides[0].evidence_verified


def test_unsupported_ai_outline_blocks_only_that_slide_and_keeps_supported_user_output():
    state = collect_context.func(
        GraphState(),
        topic="Depression",
        audience=AudienceType.SPECIALIST,
        presentation_type=PresentationType.LECTURE,
        language=Language.ENGLISH,
        duration_minutes=10,
        objective="Review treatment evidence",
    )
    presentation = create_presentation.func(validate_context.func(state)).presentation
    presentation.blueprint = Blueprint(
        title="Blueprint",
        learning_objective="Objective",
        target_number_of_slides=2,
        storytelling="Story",
        slides=[
            SlideOutline(
                slide_number=1,
                title="User conclusion",
                objective="State the user's conclusion",
                key_message="This slide was written by the user.",
                content_origin="user_authored",
            ),
            SlideOutline(
                slide_number=4,
                title="Astronomy claims",
                objective="Discuss a distant galaxy",
                key_message="Unrelated celestial statement",
            )
        ],
    )
    presentation.resources = [
        Resource(
            id="resource-1",
            filename="guideline.pdf",
            file_type=ResourceType.PDF,
            extracted_pages=[{"page": 1, "text": "Depression treatment evidence from the guideline."}],
            is_validated=True,
        )
    ]
    presentation.state.resources_validated = True
    presentation.state.workflow_status = WorkflowStatus.SLIDE_GENERATION
    result = GenerateSlidesUseCase.execute(object.__new__(GenerateSlidesUseCase), presentation)

    assert result.state.workflow_status == WorkflowStatus.AWAITING_SLIDE_RESOLUTION
    assert result.state.blocked_slide_number == 4
    assert result.state.slide_generation_error
    assert [slide.slide_number for slide in result.slides] == [1]
    assert result.slides[0].content_origin == "user_authored"
    assert [blocker.slide_number for blocker in result.state.slide_generation_blockers] == [4]

    state.presentation = result
    AuthorSlideFromBlueprintUseCase().execute(
        state,
        1,
    )
    assert state.presentation.state.workflow_status == WorkflowStatus.AWAITING_SLIDE_APPROVAL
    assert state.presentation.state.slide_generation_blockers == []
    assert [slide.slide_number for slide in state.presentation.slides] == [1, 4]
    assert state.presentation.slides[1].content_origin == "user_authored"
