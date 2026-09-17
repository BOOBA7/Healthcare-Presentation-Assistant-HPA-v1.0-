"""Owner library, deletion fences and restartable internal-export cleanup."""

import json
import hashlib
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from app.application.services.source_invalidation import invalidate_source_dependents
from app.application.services.source_screening import SourceScreening
from app.application.services.source_document import SourceDocument
from app.domain.exceptions.workflow_error import WorkflowError
from app.domain.models.resource import Resource


class SourceLifecycleRepository:
    def _init_source_lifecycle(self, connection):
        connection.execute("""CREATE TABLE IF NOT EXISTS owner_resources (
            user_id TEXT NOT NULL, resource_id TEXT NOT NULL, resource_json TEXT NOT NULL,
            original_pdf BLOB, PRIMARY KEY(user_id, resource_id))""")
        connection.execute("""CREATE TABLE IF NOT EXISTS deletion_markers (
            user_id TEXT NOT NULL, kind TEXT NOT NULL, object_id TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(user_id, kind, object_id))""")
        connection.execute("""CREATE TABLE IF NOT EXISTS source_deletion_audit (
            user_id TEXT NOT NULL, resource_id TEXT NOT NULL, actor TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(user_id, resource_id))""")
        connection.execute("""CREATE TABLE IF NOT EXISTS source_cleanup (
            filename TEXT PRIMARY KEY)""")
        connection.execute("""CREATE TABLE IF NOT EXISTS source_cleanup_state (
            id INTEGER PRIMARY KEY CHECK(id = 1), pending INTEGER NOT NULL)""")
        columns = {row[1] for row in connection.execute("PRAGMA table_info(source_cleanup_state)")}
        if "exports_root" not in columns:
            connection.execute("ALTER TABLE source_cleanup_state ADD COLUMN exports_root TEXT")
        connection.execute("INSERT OR IGNORE INTO source_cleanup_state(id, pending) VALUES (1, 0)")
        connection.execute("UPDATE source_cleanup_state SET exports_root = ? WHERE exports_root IS NULL", (str(self.legacy_exports_path.absolute()),))
        self.legacy_exports_path = Path(connection.execute("SELECT exports_root FROM source_cleanup_state WHERE id = 1").fetchone()[0])
        existing_keys = set(connection.execute("SELECT user_id, resource_id FROM owner_resources"))
        # Additive migration. No date, original, or human approval is fabricated.
        for user, project, payload in connection.execute(
            "SELECT user_id, project_id, state_json FROM project_sessions"
        ).fetchall():
            resources = {r.id: r for r in self._load_resources(connection, user, project)}
            legacy = self._deserialize_state(json.loads(payload))
            for resource in legacy.resource_library:
                if resource.extracted_text or resource.extracted_pages:
                    resources.setdefault(resource.id, resource)
            for resource in resources.values():
                if (user, resource.id) not in existing_keys:
                    self._remember_source(connection, user, resource)

    @staticmethod
    def _assert_not_deleted(connection, user, kind, identifier):
        if connection.execute(
            "SELECT 1 FROM deletion_markers WHERE user_id = ? AND kind = ? AND object_id = ?",
            (user, kind, identifier),
        ).fetchone():
            raise WorkflowError("OBJECT_DELETED", "This object was permanently deleted. Use a new identifier.")

    def _remember_source(self, connection, user, resource, *, check_only=False):
        if not connection.in_transaction:
            connection.execute("BEGIN IMMEDIATE")
        self._verify_review_authority(connection, user, resource)
        SourceScreening.resource(resource, require_text=True)
        self._assert_not_deleted(connection, user, "source", resource.id)
        previous = connection.execute(
            "SELECT original_pdf, resource_json FROM owner_resources WHERE user_id = ? AND resource_id = ?",
            (user, resource.id),
        ).fetchone()
        unchanged = (previous is not None and previous[0] == resource._original_content
                     and hashlib.sha256(resource._original_content or b"").hexdigest() == resource.metadata.original_sha256
                     and Resource.model_validate_json(previous[1]).model_dump(mode="json") == resource.model_dump(mode="json"))
        # Keeping an already accepted identical record is not a new acceptance.
        # Deletion/recovery must not depend on a currently available OCR engine.
        if resource._original_content is not None and not unchanged:
            SourceDocument.verify(resource, resource._original_content)
        if previous is not None and previous[0] != resource._original_content:
            raise WorkflowError("SOURCE_ID_CONFLICT", "This library identifier belongs to another original. Import with a new identifier.")
        if previous is None and resource._original_content is None:
            # Only already-persisted historical content may lack its original.
            payload = resource.model_dump(mode="json", exclude={"uploaded_at", "extracted_text"})
            historical = []
            for (project,) in connection.execute(
                "SELECT DISTINCT project_id FROM project_resources WHERE user_id = ? AND resource_id = ?",
                (user, resource.id),
            ).fetchall():
                historical.extend(self._load_resources(connection, user, project))
            for (state_json,) in connection.execute("SELECT state_json FROM project_sessions WHERE user_id = ?", (user,)).fetchall():
                historical.extend(self._deserialize_state(json.loads(state_json)).resource_library)
            if not any(item.id == resource.id and item._original_content is None and
                       item.model_dump(mode="json", exclude={"uploaded_at", "extracted_text"}) == payload and
                       (resource.extracted_text is None or resource.extracted_text == item.extracted_text) for item in historical):
                raise WorkflowError("SOURCE_ORIGINAL_REQUIRED", "Import the original PDF before adding a new library resource.")
        if check_only:
            return
        connection.execute(
            "INSERT OR IGNORE INTO owner_resources VALUES (?, ?, ?, ?)",
            (user, resource.id, resource.model_dump_json(), resource._original_content),
        )

    @staticmethod
    def _verify_review_authority(connection, user, resource):
        """Client JSON and direct saves cannot manufacture human confirmations."""
        from app.application.services.ocr_review import review_digest
        row = connection.execute("SELECT resource_json FROM owner_resources WHERE user_id = ? AND resource_id = ?",
                                 (user, resource.id)).fetchone()
        authoritative = Resource.model_validate_json(row[0]) if row else None
        from app.application.services.source_assets import review_digest as asset_digest
        expected_assets = authoritative.metadata.asset_reviews if authoritative else []
        if resource.metadata.asset_reviews != expected_assets:
            raise WorkflowError("ASSET_REVIEW_UNVERIFIED", "Asset reviews must use the authenticated review operation.")
        if expected_assets:
            if authoritative.metadata.assets != resource.metadata.assets or authoritative.metadata.original_sha256 != resource.metadata.original_sha256:
                raise WorkflowError("ASSET_REVIEW_UNVERIFIED", "Asset review original does not match.")
            resource._asset_review_receipt = asset_digest(resource)
        expected = authoritative.metadata.ocr_reviews if authoritative else []
        if resource.metadata.ocr_reviews != expected:
            raise WorkflowError("OCR_REVIEW_UNVERIFIED", "Extraction reviews must use the authenticated review operation. Reload this source.")
        if expected:
            if authoritative.metadata.original_sha256 != resource.metadata.original_sha256:
                raise WorkflowError("OCR_REVIEW_UNVERIFIED", "Review original does not match.")
            resource._ocr_review_receipt = review_digest(resource)

    def review_ocr(self, user, identifier, expected_revision, values, confirmed_regions):
        from app.application.services.ocr_review import apply_reviews
        from app.domain.models.source_metadata import OcrReview
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._assert_not_deleted(connection, user, "source", identifier)
            row = connection.execute("SELECT resource_json, original_pdf FROM owner_resources WHERE user_id = ? AND resource_id = ?",
                                     (user, identifier)).fetchone()
            if row is None:
                raise WorkflowError("RESOURCE_NOT_FOUND", "Source not found.")
            resource = Resource.model_validate_json(row[0])
            resource._original_content = row[1]
            self._verify_review_authority(connection, user, resource)
            SourceDocument.verify(resource, row[1])
            regions = resource.metadata.ocr_regions
            revision = max((review.batch for review in resource.metadata.ocr_reviews), default=0)
            if expected_revision != revision:
                raise WorkflowError("OCR_REVIEW_CONFLICT", "Extraction review changed. Reload before confirming.")
            if (not regions or len(values) != len(regions)
                    or sorted(confirmed_regions) != list(range(len(regions)))
                    or any(not isinstance(value, str) or not value.strip() or len(value) > 10000 for value in values)):
                raise WorkflowError("OCR_REVIEW_INCOMPLETE", "Explicitly confirm every extraction beside its original.")
            now = datetime.now(timezone.utc)
            reviews = [*resource.metadata.ocr_reviews, *[
                OcrReview(batch=revision + 1, region_index=index, original_value=region.text,
                          corrected_value=values[index].strip(), actor=user, confirmed_at=now,
                          original_sha256=resource.metadata.original_sha256)
                for index, region in enumerate(regions)
            ]]
            original = SourceDocument.read(resource.filename, row[1], resource.id)
            original.uploaded_at = resource.uploaded_at
            reviewed = apply_reviews(original, reviews)
            if resource.metadata.asset_reviews:
                from app.application.services.source_assets import apply_reviews as apply_asset_reviews
                reviewed = apply_asset_reviews(reviewed, resource.metadata.asset_reviews)
            SourceScreening.resource(reviewed, require_text=True)
            connection.execute("UPDATE owner_resources SET resource_json = ? WHERE user_id = ? AND resource_id = ?",
                               (reviewed.model_dump_json(), user, identifier))
            self._persist_source_review(connection, user, identifier, reviewed, "OCR_EXTRACTION_REVIEWED",
                                        {"review_revision": revision + 1, "region_count": len(regions), "provenance": "user_confirmed"})
        self._cleanup_if_pending()
        return {"resource_id": identifier, "review_revision": revision + 1}

    def _persist_source_review(self, connection, user, identifier, reviewed, event_type, detail):
        from app.application.services.resource_library import resource_selection
        for project, thread, payload, project_revision in connection.execute(
            "SELECT project_id, thread_id, state_json, revision FROM project_sessions WHERE user_id = ?", (user,)
        ).fetchall():
            state = self._deserialize_state(json.loads(payload))
            state.resource_library = self._load_resources(connection, user, project)
            if not any(item.id == identifier for item in state.resource_library):
                continue
            state.project_revision = project_revision
            state.resource_library = [reviewed.model_copy(deep=True) if item.id == identifier else item for item in state.resource_library]
            if state.presentation:
                state.presentation.resources = [resource_selection(reviewed) if item.id == identifier else item for item in state.presentation.resources]
            self._invalidate_project_storage(connection, user, project, state)
            self._save_project_row(connection, user, project, thread, state, None)
            self._insert_event(connection, user, project, event_type, user, {"resource_id": identifier, **detail})

    def prepare_asset_review(self, user, identifier):
        """Explicit, atomic enrichment of pre-05.2 originals; never auto-confirm."""
        from io import BytesIO
        from pptx import Presentation
        from app.domain.models.source_metadata import SourceAsset
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._assert_not_deleted(connection, user, "source", identifier)
            row = connection.execute("SELECT resource_json, original_pdf FROM owner_resources WHERE user_id = ? AND resource_id = ?",
                                     (user, identifier)).fetchone()
            if row is None:
                raise WorkflowError("RESOURCE_NOT_FOUND", "Source not found.")
            resource = Resource.model_validate_json(row[0])
            resource._original_content = row[1]
            self._verify_review_authority(connection, user, resource)
            if resource.metadata.assets and all(asset.content_base64 for asset in resource.metadata.assets):
                SourceDocument.verify(resource, row[1])
                return {"asset_count": len(resource.metadata.assets)}
            if resource.metadata.asset_reviews or resource.file_type.value not in ('pdf', 'pptx'):
                raise WorkflowError("ASSET_INTEGRITY_FAILED", "Replace this source with its verified original.")
            fresh = SourceDocument.read(resource.filename, row[1], identifier)
            if resource.file_type.value == 'pptx':
                legacy = []
                shapes = [shape for slide in Presentation(BytesIO(row[1])).slides for shape in slide.shapes
                          if shape.has_table or shape.shape_type == 13]
                if len(shapes) != len(fresh.metadata.assets):
                    raise WorkflowError("ASSET_INTEGRITY_FAILED", "Original inventory does not match.")
                for shape, asset in zip(shapes, fresh.metadata.assets):
                    legacy.append(SourceAsset(id=f'slide-{asset.location.number}-shape-{shape.shape_id}', kind=asset.kind,
                                              media_type=None if asset.kind == 'table' else asset.media_type,
                                              location=asset.location))
                if resource.metadata.assets != legacy:
                    raise WorkflowError("ASSET_INTEGRITY_FAILED", "Original inventory does not match.")
            elif resource.metadata.assets:
                raise WorkflowError("ASSET_INTEGRITY_FAILED", "Original inventory does not match.")
            candidate = resource.model_copy(deep=True)
            candidate.metadata.assets = fresh.metadata.assets
            # Re-derives every other field, including any authenticated OCR review.
            SourceDocument.verify(candidate, row[1])
            SourceScreening.resource(candidate, require_text=True)
            if not fresh.metadata.assets:
                return {"asset_count": 0}
            connection.execute("UPDATE owner_resources SET resource_json = ? WHERE user_id = ? AND resource_id = ?",
                               (candidate.model_dump_json(), user, identifier))
            self._persist_source_review(connection, user, identifier, candidate, "SOURCE_ASSETS_RETAINED",
                                        {"asset_count": len(candidate.metadata.assets)})
        self._cleanup_if_pending()
        return {"asset_count": len(candidate.metadata.assets)}

    def review_assets(self, user, identifier, expected_revision, values, confirmed_assets):
        from app.application.services.source_assets import FIELDS, apply_reviews
        from app.application.services.ocr_review import apply_reviews as apply_ocr_reviews
        from app.domain.models.source_metadata import AssetReview
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            self._assert_not_deleted(connection, user, "source", identifier)
            row = connection.execute("SELECT resource_json, original_pdf FROM owner_resources WHERE user_id = ? AND resource_id = ?",
                                     (user, identifier)).fetchone()
            if row is None:
                raise WorkflowError("RESOURCE_NOT_FOUND", "Source not found.")
            resource = Resource.model_validate_json(row[0])
            resource._original_content = row[1]
            self._verify_review_authority(connection, user, resource)
            SourceDocument.verify(resource, row[1])
            assets = resource.metadata.assets
            revision = max((review.revision for review in resource.metadata.asset_reviews), default=0)
            if type(expected_revision) is not int or expected_revision != revision:
                raise WorkflowError("ASSET_REVIEW_CONFLICT", "Asset review changed. Reload before confirming.")
            if (not assets or len(values) != len(assets) or sorted(confirmed_assets) != sorted(a.id for a in assets)
                    or any(not isinstance(value, dict) or set(value) != set(FIELDS)
                           or any(item is not None and (not isinstance(item, str) or len(item) > 10000) for item in value.values())
                           for value in values)):
                raise WorkflowError("ASSET_REVIEW_INCOMPLETE", "Explicitly confirm every item beside its original.")
            now = datetime.now(timezone.utc)
            reviews = [*resource.metadata.asset_reviews, *[
                AssetReview(revision=revision + 1, asset_id=asset.id, original_sha256=resource.metadata.original_sha256,
                            asset_sha256=asset.sha256, initial={name: getattr(asset, name) for name in FIELDS},
                            corrected=values[index], actor=user, confirmed_at=now)
                for index, asset in enumerate(assets)
            ]]
            original = SourceDocument.read(resource.filename, row[1], resource.id)
            original.uploaded_at = resource.uploaded_at
            if resource.metadata.ocr_reviews:
                original = apply_ocr_reviews(original, resource.metadata.ocr_reviews)
            reviewed = apply_reviews(original, reviews)
            SourceScreening.resource(reviewed, require_text=True)
            connection.execute("UPDATE owner_resources SET resource_json = ? WHERE user_id = ? AND resource_id = ?",
                               (reviewed.model_dump_json(), user, identifier))
            self._persist_source_review(connection, user, identifier, reviewed, "SOURCE_ASSETS_REVIEWED",
                                        {"review_revision": revision + 1, "asset_count": len(assets), "provenance": "user_confirmed"})
            # The append-only review rows inside owner_resources are also the
            # private audit when the source is no longer attached to any Project.
        self._cleanup_if_pending()
        return {"resource_id": identifier, "review_revision": revision + 1}

    def list_library_resources(self, user):
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT resource_json, original_pdf FROM owner_resources WHERE user_id = ? ORDER BY resource_id", (user,)
            ).fetchall()
        result = []
        for payload, content in rows:
            resource = Resource.model_validate_json(payload)
            resource._original_content = content
            if resource.metadata.ocr_reviews:
                from app.application.services.ocr_review import review_digest
                resource._ocr_review_receipt = review_digest(resource)
            if resource.metadata.asset_reviews:
                from app.application.services.source_assets import review_digest as asset_digest
                resource._asset_review_receipt = asset_digest(resource)
            result.append(resource)
        return result

    def library_resource(self, user, identifier):
        return next((r for r in self.list_library_resources(user) if r.id == identifier), None)

    def _queue_project_exports(self, connection, user, project, state):
        # Only application-created filenames inside the internal exports root.
        names = set()
        for (payload,) in connection.execute(
            "SELECT payload_json FROM project_events WHERE user_id = ? AND project_id = ? AND event_type = 'PRESENTATION_EXPORTED'",
            (user, project),
        ):
            name = json.loads(payload).get("filename", "")
            if re.fullmatch(r"[0-9a-fA-F-]{36}-[0-9a-f]{8}\.pptx", name):
                names.add(name)
        if state.presentation and re.fullmatch(r"[0-9a-fA-F-]{36}", state.presentation.id):
            # Includes old exports interrupted between file creation and audit.
            prefix = state.presentation.id
            names.update(p.name for p in self.legacy_exports_path.glob(f"{prefix}-*.pptx")
                         if re.fullmatch(re.escape(prefix) + r"-[0-9a-f]{8}\.pptx", p.name))
        if names:
            for (other_payload,) in connection.execute("SELECT state_json FROM project_sessions WHERE user_id != ?", (user,)):
                other_id = (json.loads(other_payload).get("presentation") or {}).get("id")
                if other_id and any(name.startswith(other_id + "-") for name in names):
                    raise WorkflowError("EXPORT_OWNERSHIP_AMBIGUOUS", "A legacy presentation identifier is shared across owners; internal export ownership must be resolved before deletion.")
        connection.executemany("INSERT OR IGNORE INTO source_cleanup VALUES (?)", [(name,) for name in names])
        connection.execute("UPDATE source_cleanup_state SET pending = pending + 1 WHERE id = 1")

    def _invalidate_project_storage(self, connection, user, project, state):
        self._queue_project_exports(connection, user, project, state)
        invalidate_source_dependents(state)
        connection.execute("""UPDATE project_jobs SET status = 'cancelled', stage = 'source_changed',
            result_json = NULL, error_message = NULL, updated_at = CURRENT_TIMESTAMP
            WHERE user_id = ? AND project_id = ?""", (user, project))
        # Preserve who/when/type, remove potentially quoted text in old payloads.
        connection.execute("""UPDATE project_events SET payload_json = '{"redacted_after_source_change": true}'
            WHERE user_id = ? AND project_id = ?""", (user, project))

    def permanently_delete_source(self, user, identifier, actor):
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            exists = connection.execute("SELECT 1 FROM owner_resources WHERE user_id = ? AND resource_id = ?", (user, identifier)).fetchone()
            marker = connection.execute("SELECT 1 FROM deletion_markers WHERE user_id = ? AND kind = 'source' AND object_id = ?", (user, identifier)).fetchone()
            if not exists and not marker:
                return None
            connection.execute("INSERT OR IGNORE INTO deletion_markers(user_id, kind, object_id) VALUES (?, 'source', ?)", (user, identifier))
            for project, thread, payload, revision in connection.execute(
                "SELECT project_id, thread_id, state_json, revision FROM project_sessions WHERE user_id = ?", (user,)
            ).fetchall():
                state = self._deserialize_state(json.loads(payload))
                normalized = self._load_resources(connection, user, project)
                if normalized:
                    state.resource_library = normalized
                ids = {r.id for r in state.resource_library}
                if state.presentation:
                    ids.update(r.id for r in state.presentation.resources)
                if identifier not in ids:
                    continue
                state.project_revision = revision
                state.resource_library = [r for r in state.resource_library if r.id != identifier]
                if state.presentation:
                    state.presentation.resources = [r for r in state.presentation.resources if r.id != identifier]
                self._invalidate_project_storage(connection, user, project, state)
                self._save_project_row(connection, user, project, thread, state, None)
                self._insert_event(connection, user, project, "SOURCE_PERMANENTLY_DELETED", actor, {"resource_id": identifier})
            for table in ("project_resource_chunks", "project_resource_pages", "project_resources", "owner_resources"):
                connection.execute(f"DELETE FROM {table} WHERE user_id = ? AND resource_id = ?", (user, identifier))
            connection.execute("INSERT OR IGNORE INTO source_deletion_audit(user_id, resource_id, actor) VALUES (?, ?, ?)", (user, identifier, actor))
            connection.execute("UPDATE source_cleanup_state SET pending = pending + 1 WHERE id = 1")
        return {"resource_id": identifier, "deleted": True, "cleanup_pending": not self.finish_source_cleanup()}

    def source_cleanup_pending(self):
        with self._connect() as connection:
            return bool(connection.execute("SELECT pending FROM source_cleanup_state WHERE id = 1").fetchone()[0])

    def finish_source_cleanup(self):
        """Idempotent after commit and on restart; never follow export symlinks."""
        try:
            with self._connect() as connection:
                connection.execute("BEGIN IMMEDIATE")
                generation = connection.execute("SELECT pending FROM source_cleanup_state WHERE id = 1").fetchone()[0]
                for (filename,) in connection.execute("SELECT filename FROM source_cleanup").fetchall():
                    if not re.fullmatch(r"[0-9a-fA-F-]{36}-[0-9a-f]{8}\.pptx", filename):
                        return False
                    if self.legacy_exports_path.is_symlink():
                        return False
                    (self.legacy_exports_path / filename).unlink(missing_ok=True)
                    connection.execute("DELETE FROM source_cleanup WHERE filename = ?", (filename,))
            with self._connect() as connection:
                busy, _, _ = connection.execute("PRAGMA wal_checkpoint(TRUNCATE)").fetchone()
                if busy:
                    return False
                connection.execute("UPDATE source_cleanup_state SET pending = 0 WHERE id = 1 AND pending = ?", (generation,))
            return True
        except (OSError, sqlite3.Error):
            return False

    def _delete_project_rows(self, connection, user, project):
        row = connection.execute("SELECT state_json FROM project_sessions WHERE user_id = ? AND project_id = ?", (user, project)).fetchone()
        if row is None:
            return False
        self._queue_project_exports(connection, user, project, self._deserialize_state(json.loads(row[0])))
        connection.execute("INSERT OR IGNORE INTO deletion_markers(user_id, kind, object_id) VALUES (?, 'project', ?)", (user, project))
        for table in ("project_resource_chunks", "project_resource_pages", "project_resources", "project_jobs", "project_events", "project_sessions"):
            connection.execute(f"DELETE FROM {table} WHERE user_id = ? AND project_id = ?", (user, project))
        return True

    def delete_project(self, user_id, project_id):
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            result = self._delete_project_rows(connection, user_id, project_id)
        self.finish_source_cleanup()
        return result
