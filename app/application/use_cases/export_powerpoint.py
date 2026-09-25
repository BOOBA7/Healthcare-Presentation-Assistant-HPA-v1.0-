from app.application.services.pptx_roles import reference_warnings, source_warning
from pathlib import Path
from io import BytesIO
from uuid import uuid4

from pptx import Presentation as PowerPoint
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.application.validators.evidence_provenance_validator import EvidenceProvenanceValidator
from app.domain.enums.presentation_theme import PresentationColour, PresentationTheme
from app.domain.models.presentation import Presentation
from app.domain.models.resource import Resource
from app.application.services.workflow_policy import WorkflowPolicy


THEMES: dict[PresentationTheme, dict[str, tuple[int, int, int]]] = {
    PresentationTheme.CLINICAL: {"primary": (16, 89, 104), "accent": (49, 161, 153), "paper": (245, 250, 250), "ink": (20, 43, 54)},
    PresentationTheme.ACADEMIC: {"primary": (30, 55, 98), "accent": (194, 145, 62), "paper": (248, 248, 252), "ink": (27, 40, 72)},
    PresentationTheme.EXECUTIVE: {"primary": (36, 44, 61), "accent": (209, 113, 75), "paper": (250, 249, 247), "ink": (36, 44, 61)},
    PresentationTheme.MIDNIGHT: {"primary": (15, 23, 42), "accent": (56, 189, 248), "paper": (241, 245, 249), "ink": (15, 23, 42)},
}

COLOURS: dict[PresentationColour, dict[str, tuple[int, int, int]]] = {
    PresentationColour.TEAL: THEMES[PresentationTheme.CLINICAL],
    PresentationColour.BLUE: THEMES[PresentationTheme.ACADEMIC],
    PresentationColour.WARM: THEMES[PresentationTheme.EXECUTIVE],
}


class ExportPowerPointUseCase:
    """Export a human-approved presentation in a professional visual theme."""

    @staticmethod
    def require_eligible(presentation: Presentation, resources: list[Resource]) -> None:
        """Apply the single deterministic gate shared by preview and exports."""
        from app.application.services.source_date_policy import SourceDatePolicy

        SourceDatePolicy.require_all(resources)
        if not presentation.slides:
            raise ValueError("Generate presentation slides before export.")
        if not presentation.state.presentation_validated:
            raise ValueError("Human approval of the final presentation is required before export.")
        WorkflowPolicy.require_export_eligible(presentation)
        if not resources or not presentation.state.resources_validated:
            raise ValueError("Validated user resources are required before export.")
        if not all(resource.is_validated for resource in resources):
            raise ValueError("Every resource must be validated by the user before export.")
        if presentation.agenda is None or not presentation.agenda.is_validated:
            raise ValueError("A user-approved agenda is required before export.")
        if presentation.blueprint is None or not presentation.state.blueprint_validated:
            raise ValueError("A user-approved Blueprint is required before export.")
        EvidenceProvenanceValidator().validate_presentation(
            [slide.model_copy(deep=True) for slide in presentation.slides], resources
        )

    def execute(
        self,
        presentation: Presentation,
        output_dir: Path | BytesIO,
        graphic_source: Resource | None = None,
        resources: list[Resource] | None = None,
    ) -> Path | BytesIO:
        resources = resources if resources is not None else presentation.resources
        self.require_eligible(presentation, resources)

        if isinstance(output_dir, Path):
            output_dir.mkdir(parents=True, exist_ok=True)
        if presentation.custom_template_id or graphic_source is not None:
            from app.application.services.source_document import SourceDocument
            if (not isinstance(graphic_source, Resource)
                    or graphic_source.id != presentation.custom_template_id
                    or graphic_source.file_type.value != "pptx"
                    or graphic_source._original_content is None):
                raise ValueError("Select a screened local PowerPoint through the explicit style action.")
            SourceDocument.verify(graphic_source, graphic_source._original_content)
        deck = PowerPoint(BytesIO(graphic_source._original_content)) if graphic_source else PowerPoint()
        if graphic_source:
            self._remove_template_slides(deck)
        else:
            deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        palette = THEMES[presentation.theme] if presentation.colour == PresentationColour.THEME else COLOURS[presentation.colour]
        citation_markers = self._citation_markers(presentation)

        self._add_title_slide(deck, presentation, palette)
        self._add_agenda_slide(deck, presentation, palette)
        for slide in presentation.slides:
            self._add_content_slide(deck, slide, palette, resources, citation_markers)
            claim_references = [{"resource_id": link.resource_id} for link in slide.evidence_links]
            for warning in reference_warnings(slide.reference_details + claim_references, resources):
                self._textbox(deck.slides[-1], warning, 0.8, 6.82, 11.7, 0.25, size=9, color=palette["ink"])
            if slide.speaker_notes:
                deck.slides[-1].notes_slide.notes_text_frame.text = slide.speaker_notes
        self._add_resources_slide(deck, presentation, resources, palette)

        path = output_dir / f"{presentation.id}-{uuid4().hex[:8]}.pptx" if isinstance(output_dir, Path) else output_dir
        deck.save(path)
        return path

    @staticmethod
    def _remove_template_slides(deck: PowerPoint) -> None:
        """Keep the uploaded template's masters/layouts while removing sample slides."""
        while deck.slides:
            relationship_id = deck.slides._sldIdLst[0].rId
            deck.part.drop_rel(relationship_id)
            del deck.slides._sldIdLst[0]

    @staticmethod
    def _blank_layout(deck: PowerPoint):
        return deck.slide_layouts[6] if len(deck.slide_layouts) > 6 else deck.slide_layouts[-1]

    @staticmethod
    def _rgb(value: tuple[int, int, int]) -> RGBColor:
        return RGBColor(*value)

    def _background(self, slide, color: tuple[int, int, int]) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(color)

    def _textbox(self, slide, text: str, left: float, top: float, width: float, height: float, *, size: int, color: tuple[int, int, int], bold: bool = False, align=PP_ALIGN.LEFT):
        frame = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = align
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = self._rgb(color)
        return frame

    def _add_title_slide(self, deck: PowerPoint, presentation: Presentation, palette: dict[str, tuple[int, int, int]]) -> None:
        slide = deck.slides.add_slide(self._blank_layout(deck))
        self._background(slide, palette["primary"])
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.25), Inches(7.5))
        accent.fill.solid()
        accent.fill.fore_color.rgb = self._rgb(palette["accent"])
        accent.line.fill.background()
        self._textbox(slide, "HEALTHCARE PRESENTATION", 0.8, 1.25, 10.8, 0.35, size=12, color=palette["accent"], bold=True)
        self._textbox(slide, presentation.title, 0.8, 1.8, 11.5, 1.6, size=34, color=(255, 255, 255), bold=True)
        self._textbox(slide, presentation.context.objective, 0.8, 3.75, 10.7, 0.62, size=17, color=(220, 235, 235))
        details = self._title_slide_details(presentation)
        for index, detail in enumerate(details):
            self._textbox(slide, detail, 0.8, 4.55 + index * 0.34, 11.4, 0.28, size=11, color=(220, 235, 235))
        self._textbox(slide, f"{presentation.context.duration_minutes} min  ·  {presentation.context.audience.value}", 0.8, 6.5, 10.0, 0.3, size=11, color=(220, 235, 235))

    @staticmethod
    def _title_slide_details(presentation: Presentation) -> list[str]:
        context = presentation.context
        details: list[str] = []
        presenter = " · ".join(item for item in (context.presenter_name, context.presenter_title) if item)
        location = " · ".join(item for item in (context.venue, context.presentation_date) if item)
        for item in (presenter, context.organization, context.event_name, location):
            if item:
                details.append(item)
        return details[:4]

    def _add_agenda_slide(self, deck: PowerPoint, presentation: Presentation, palette: dict[str, tuple[int, int, int]]) -> None:
        slide = deck.slides.add_slide(self._blank_layout(deck))
        self._background(slide, palette["paper"])
        self._header(slide, self._label(presentation, "Agenda", "Agenda", "جدول الأعمال"), palette)
        for index, item in enumerate(presentation.agenda.items):
            top = 1.55 + index * 0.78
            badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(top), Inches(0.38), Inches(0.38))
            badge.fill.solid()
            badge.fill.fore_color.rgb = self._rgb(palette["accent"])
            badge.line.fill.background()
            self._textbox(slide, str(index + 1), 0.8, top + 0.02, 0.38, 0.25, size=10, color=(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
            self._textbox(slide, item, 1.4, top - 0.02, 10.7, 0.45, size=20, color=palette["ink"], bold=True)
        self._footer(slide, palette)

    def _add_content_slide(self, deck: PowerPoint, source_slide, palette: dict[str, tuple[int, int, int]], resources: list[Resource], citation_markers: dict[str, int] | None = None) -> None:
        slide = deck.slides.add_slide(self._blank_layout(deck))
        self._background(slide, palette["paper"])
        self._header(slide, source_slide.title, palette)
        visual = source_slide.visual
        text_left = source_slide.layout != "visual_left_text_right"
        text_width = 5.45 if visual else 11.5
        text_x = 0.85 if text_left else 7.0
        if source_slide.layout != "visual_focus":
            bullet_frame = slide.shapes.add_textbox(Inches(text_x), Inches(1.65), Inches(text_width), Inches(3.85)).text_frame
            bullet_frame.clear()
            for index, message in enumerate(source_slide.key_messages or [source_slide.content]):
                paragraph = bullet_frame.paragraphs[0] if index == 0 else bullet_frame.add_paragraph()
                paragraph.text = message
                paragraph.level = 0
                paragraph.font.size = Pt(22)
                paragraph.font.color.rgb = self._rgb(palette["ink"])
                paragraph.space_after = Pt(13)
        if visual:
            visual_x = 6.75 if text_left else 0.75
            if source_slide.layout == "visual_focus":
                visual_x, visual_width = 2.05, 9.25
            else:
                visual_width = 5.75
            self._add_visual(slide, visual, resources, palette, visual_x, 1.55, visual_width, 3.9)
        if source_slide.content_origin == "ai_generated":
            references = [{
                "title": link.resource_title, "resource_id": link.resource_id,
                "location_kind": link.location_kind, "page": link.location_number,
                "evidence_excerpt": link.exact_passage,
            } for link in source_slide.evidence_links
                if link.provenance_verified and link.semantic_review == "approved"]
            references.extend(source_slide.reference_details)
            self._add_evidence_box(slide, references, palette, citation_markers or {})
        else:
            self._add_authorship_box(slide, source_slide, palette)
        self._footer(slide, palette, self._content_origin_label(source_slide))

    def _add_visual(self, slide, visual, resources, palette, left, top, width, height):
        if visual.kind == "image":
            from app.application.services.source_assets import asset_content, is_reviewed
            resource = next((item for item in resources if item.id == visual.resource_id), None)
            if resource is None or not is_reviewed(resource):
                raise ValueError("A reviewed supplied image is required for export.")
            content, media_type = asset_content(resource, visual.asset_id)
            if not media_type or not media_type.startswith("image/"):
                raise ValueError("The selected supplied asset is not an image.")
            slide.shapes.add_picture(BytesIO(content), Inches(left), Inches(top), Inches(width), Inches(height))
            return
        if visual.kind == "table":
            shape = slide.shapes.add_table(len(visual.rows) + 1, len(visual.columns), Inches(left), Inches(top), Inches(width), Inches(height))
            for column, label in enumerate(visual.columns):
                shape.table.cell(0, column).text = label
            for row_index, row in enumerate(visual.rows, 1):
                for column, value in enumerate(row):
                    shape.table.cell(row_index, column).text = value
            return
        if visual.kind in {"bar_chart", "line_chart"}:
            data = ChartData()
            data.categories = visual.categories
            for name, values in visual.series.items():
                data.add_series(name, values)
            chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if visual.kind == "bar_chart" else XL_CHART_TYPE.LINE_MARKERS
            slide.shapes.add_chart(chart_type, Inches(left), Inches(top), Inches(width), Inches(height), data)
            return
        gap = width / max(len(visual.nodes), 1)
        for index, label in enumerate(visual.nodes):
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + index * gap), Inches(top + 1.1), Inches(max(gap - .18, .5)), Inches(1.05))
            box.fill.solid()
            box.fill.fore_color.rgb = self._rgb(palette["accent"])
            box.text_frame.text = label
            if index:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left + index * gap - .18), Inches(top + 1.43), Inches(.22), Inches(.3))
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = self._rgb(palette["primary"])

    def _header(self, slide, title: str, palette: dict[str, tuple[int, int, int]]) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self._rgb(palette["primary"])
        bar.line.fill.background()
        self._textbox(slide, title, 0.8, 0.29, 11.9, 0.47, size=26, color=(255, 255, 255), bold=True)

    def _footer(self, slide, palette: dict[str, tuple[int, int, int]], provenance: str | None = None) -> None:
        text = "HPA  ·  Human-reviewed scientific presentation"
        if provenance:
            text = f"{text}  ·  {provenance}"
        self._textbox(slide, text, 0.8, 7.08, 11.7, 0.18, size=8, color=palette["ink"])

    def _add_evidence_box(self, slide, references: list[dict[str, object]], palette: dict[str, tuple[int, int, int]], citation_markers: dict[str, int]) -> None:
        if not references:
            return
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.65), Inches(11.7), Inches(1.15))
        box.fill.solid()
        box.fill.fore_color.rgb = self._rgb((230, 240, 239))
        box.line.color.rgb = self._rgb(palette["accent"])
        frame = box.text_frame
        frame.clear()
        frame.margin_left = Inches(0.14)
        frame.margin_top = Inches(0.08)
        for index, reference in enumerate(references):
            title = str(reference.get("title") or "Source")
            excerpt = " ".join(str(reference.get("evidence_excerpt") or "").split())
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            location = "slide" if reference.get("location_kind") == "slide" else "p."
            marker = citation_markers.get(str(reference.get("resource_id")), index + 1)
            paragraph.text = f"[{marker}] {title} — {reference.get('resource_id')}, {location} {reference.get('page')}: « {excerpt} »"
            paragraph.font.size = Pt(7)
            paragraph.font.color.rgb = self._rgb(palette["ink"])

    @staticmethod
    def _citation_markers(presentation: Presentation) -> dict[str, int]:
        cited_ids: list[str] = []
        for source_slide in presentation.slides:
            references = [*source_slide.reference_details, *source_slide.evidence_links]
            if source_slide.speaker_note is not None:
                references.extend(source_slide.speaker_note.evidence_links)
            for reference in references:
                resource_id = reference.get("resource_id") if isinstance(reference, dict) else reference.resource_id
                if resource_id and resource_id not in cited_ids:
                    cited_ids.append(resource_id)
        return {resource_id: index + 1 for index, resource_id in enumerate(cited_ids)}

    def _add_authorship_box(self, slide, source_slide, palette: dict[str, tuple[int, int, int]]) -> None:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.82), Inches(11.7), Inches(0.72))
        box.fill.solid()
        box.fill.fore_color.rgb = self._rgb((255, 247, 237))
        box.line.color.rgb = self._rgb(palette["accent"])
        message = (
            "User-authored content · Human-approved"
            if source_slide.content_origin == "user_authored"
            else "User-edited content · Human-approved"
        )
        self._textbox(slide, message, 1.0, 6.04, 11.2, 0.22, size=9, color=palette["ink"], bold=True)

    @staticmethod
    def _content_origin_label(source_slide) -> str:
        # Kept short because it appears on every content slide's footer.
        if source_slide.content_origin == "user_authored":
            return "User-authored · Human-approved"
        if source_slide.content_origin == "user_edited":
            return "User-edited · Human-approved"
        return "AI-generated · Human-approved"

    def _add_resources_slide(
        self,
        deck: PowerPoint,
        presentation: Presentation,
        resources: list[Resource],
        palette: dict[str, tuple[int, int, int]],
    ) -> None:
        # One resource is deliberately a small visual block rather than one
        # long sentence. IDs remain available for audit, but do not compete
        # with the document title during normal reading.
        batches = [resources[index:index + 4] for index in range(0, len(resources), 4)]
        citation_markers = self._citation_markers(presentation)
        origins = {
            "ai_generated": sum(item.content_origin == "ai_generated" for item in presentation.slides),
            "user_edited": sum(item.content_origin == "user_edited" for item in presentation.slides),
            "user_authored": sum(item.content_origin == "user_authored" for item in presentation.slides),
        }
        title = self._label(presentation, "Resources and validation", "Ressources et validation", "المصادر والاعتماد")
        validated = self._label(
            presentation,
            "All resources below were validated by the user.",
            "Toutes les ressources ci-dessous ont été validées par l’utilisateur.",
            "تم اعتماد جميع المصادر التالية من قبل المستخدم.",
        )
        source_label = self._label(presentation, "Organisation", "Organisme", "الجهة")
        audit_label = self._label(presentation, "Audit ID", "Identifiant d’audit", "معرّف التدقيق")

        for batch_number, batch in enumerate(batches, start=1):
            slide = deck.slides.add_slide(self._blank_layout(deck))
            self._background(slide, palette["paper"])
            suffix = f" ({batch_number}/{len(batches)})" if len(batches) > 1 else ""
            self._header(slide, f"{title}{suffix}", palette)

            if batch_number == 1:
                self._textbox(slide, validated, 0.85, 1.38, 11.5, 0.38, size=15, color=palette["ink"], bold=True)
                provenance = (
                    f"Content provenance: {origins['ai_generated']} AI-generated, "
                    f"{origins['user_edited']} user-edited, {origins['user_authored']} user-authored."
                )
                self._textbox(slide, provenance, 0.85, 1.86, 11.5, 0.3, size=9, color=palette["ink"])
                top = 2.38
            else:
                top = 1.5

            for index, resource in enumerate(batch):
                resource_number = (batch_number - 1) * 4 + index + 1
                spacing = 1.10 if any(source_warning(item) for item in batch) else 0.98
                item_top = top + index * spacing
                self._textbox(
                    slide,
                    f"{f'[{citation_markers[resource.id]}] ' if resource.id in citation_markers else ''}{resource_number}. {resource.title or resource.filename}",
                    0.95,
                    item_top,
                    11.1,
                    0.42,
                    size=13,
                    color=palette["ink"],
                    bold=True,
                )
                if source_warning(resource):
                    self._textbox(slide, source_warning(resource), 1.18, item_top + 0.86, 11.1, 0.22, size=8, color=palette["ink"])
                if resource.source:
                    self._textbox(
                        slide,
                        f"{source_label}: {resource.source}",
                        1.18,
                        item_top + 0.43,
                        10.9,
                        0.2,
                        size=9,
                        color=palette["ink"],
                    )
                self._textbox(
                    slide,
                    f"{audit_label}: {resource.id}",
                    1.18,
                    item_top + 0.66,
                    10.9,
                    0.18,
                    size=8,
                    color=palette["ink"],
                )
            self._footer(slide, palette)

    @staticmethod
    def _label(presentation: Presentation, english: str, french: str, arabic: str) -> str:
        language = presentation.context.language.value
        return french if language == "French" else arabic if language == "Arabic" else english
