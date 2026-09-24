"""Deterministic retained derivatives, source-attributed fields and review gates.

Bytes live in the existing private source record and its atomic lifecycle; no
filesystem cache or temporary copy is created. Never send these bytes to a model.
"""
import base64
import hashlib
import json
import re
import time
from io import BytesIO

import fitz
from pptx import Presentation

from app.application.services.raster_privacy import screen_raster_text
from app.domain.exceptions.workflow_error import WorkflowError
from app.domain.models.source_metadata import AssetField, SourceAsset, SourceLocation

FIELDS = ('caption', 'rights', 'author', 'organisation', 'provenance')
MAX_DERIVED_BYTES = 10 * 1024 * 1024
MAX_ASSETS = 200
MAX_SECONDS = 60


def failure():
    return WorkflowError('ASSET_INTEGRITY_FAILED', 'An extracted item does not match its verified original.')


def make_asset(content, original, locator, kind, location, media_type, candidates=None):
    if len(content) > MAX_DERIVED_BYTES:
        raise failure()
    digest = hashlib.sha256(content).hexdigest()
    original_hash = hashlib.sha256(original).hexdigest()
    identity = json.dumps([original_hash, locator, kind, location.model_dump(mode='json'), digest], sort_keys=True)
    fields = {name: AssetField() for name in FIELDS}
    for name, entries in (candidates or {}).items():
        values, origins = zip(*entries)
        for value in values:
            screen_raster_text(value)
        fields[name] = AssetField(status='present' if len(set(values)) == 1 else 'contradictory',
                                  values=list(values), origins=list(origins))
    return SourceAsset(id=hashlib.sha256(identity.encode()).hexdigest(), kind=kind, location=location,
                       media_type=media_type, original_sha256=original_hash, sha256=digest,
                       content_base64=base64.b64encode(content).decode(), fields=fields,
                       **{name: field.values[0] if field.status == 'present' else None for name, field in fields.items()})


def table_bytes(rows):
    for row in rows:
        for value in row:
            screen_raster_text(value or '')
    return json.dumps(rows, ensure_ascii=False, separators=(',', ':')).encode()


def unattributed_fields(asset, text):
    """Nearby labels cannot establish ownership of a particular figure/table."""
    for name in FIELDS:
        matches = re.findall(r'^' + name + r':\s*(.+)$', text, re.MULTILINE | re.IGNORECASE)
        if matches and asset.fields[name].status == 'missing':
            for value in matches:
                screen_raster_text(value)
            asset.fields[name] = AssetField(status='not_attributable' if len(set(matches)) == 1 else 'ambiguous',
                                            values=matches, origins=['nearby_text_unattributed'] * len(matches))
    return asset


def image_candidates(metadata):
    result = {}
    for key, field in [('Description', 'caption'), ('Copyright', 'rights'), ('Author', 'author')]:
        if metadata.get(key):
            result.setdefault(field, []).append((metadata[key], 'embedded_image:' + key))
    from app.application.services.source_date_policy import SourceDatePolicy
    root = SourceDatePolicy.xml_root(metadata.get('XML:com.adobe.xmp'))
    if root is not None:
        for key, field in [('description', 'caption'), ('rights', 'rights'), ('creator', 'author'), ('publisher', 'organisation'), ('source', 'provenance')]:
            for element in root.iter('{http://purl.org/dc/elements/1.1/}' + key):
                value = ' '.join(text.strip() for text in element.itertext() if text.strip())
                if value:
                    result.setdefault(field, []).append((value, 'embedded_image:xmp:dc:' + key))
    return result


def check_budget(assets, started):
    if (len(assets) > MAX_ASSETS or sum(len(asset.content_base64 or '') for asset in assets) > MAX_DERIVED_BYTES
            or time.monotonic() - started > MAX_SECONDS):
        raise failure()


def pptx_assets(content):
    """Only called after the full package, metadata and pixels passed screening."""
    from app.application.services.raster_document import RasterDocument
    from app.application.services.pptx_document import P
    presentation = Presentation(BytesIO(content))
    started = time.monotonic()
    result = []
    for number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if not (shape.has_table or shape._element.tag == f'{{{P}}}pic'):
                continue
            box = (shape.left / presentation.slide_width, shape.top / presentation.slide_height,
                   (shape.left + shape.width) / presentation.slide_width,
                   (shape.top + shape.height) / presentation.slide_height)
            candidates = {}
            if shape.has_table:
                data = table_bytes([[cell.text for cell in row.cells] for row in shape.table.rows])
                kind, media = 'table', 'application/json'
            else:
                data = shape.image.blob
                _, _, metadata, _ = RasterDocument.image(data)
                candidates = image_candidates(metadata)
                kind, media = 'image', shape.image.content_type
            properties = shape._element.find(f'.//{{{P}}}cNvPr')
            for key in ('descr', 'title'):
                if properties is not None and properties.get(key) and not re.fullmatch(r'image\.(?:png|jpe?g)', properties.get(key), re.IGNORECASE):
                    candidates.setdefault('caption', []).append((properties.get(key), f'shape:{shape.shape_id}:{key}'))
            asset = make_asset(data, content, f'slide:{number}:shape:{shape.shape_id}', kind,
                               SourceLocation(kind='slide', number=number, region=box), media, candidates)
            nearby = '\n'.join(item.text for item in slide.shapes if item.has_text_frame)
            result.append(unattributed_fields(asset, nearby))
            check_budget(result, started)
    return result


def aligned_tables(page):
    """Conservative native-text candidates: repeated columns separated by 40pt.

    No guessed cell merges, grid/vector interpretation or multiline cells.
    The original word boxes, not reading/file order, define each candidate.
    """
    lines = []
    for word in sorted(page.get_text('words'), key=lambda w: (w[1], w[0])):
        if not lines or abs(word[1] - lines[-1][0][1]) > 2:
            lines.append([])
        lines[-1].append(word)
    groups, current, previous_starts = [], [], None
    for line in lines:
        cells = []
        for word in sorted(line, key=lambda w: w[0]):
            if not cells or word[0] - cells[-1][-1][2] >= 40:
                cells.append([])
            cells[-1].append(word)
        starts = [cell[0][0] for cell in cells]
        compatible = (len(cells) >= 2 and previous_starts is not None and len(starts) == len(previous_starts)
                      and all(abs(a - b) <= 2 for a, b in zip(starts, previous_starts))
                      and line[0][1] - current[-1][0][0][1] <= 40)
        if not compatible:
            if len(current) >= 2:
                groups.append(current)
            current = []
        if len(cells) >= 2:
            current.append(cells)
            previous_starts = starts
        else:
            previous_starts = None
    if len(current) >= 2:
        groups.append(current)
    for group in groups:
        words = [word for row in group for cell in row for word in cell]
        rectangle = (min(w[0] for w in words), min(w[1] for w in words),
                     max(w[2] for w in words), max(w[3] for w in words))
        yield [[' '.join(word[4] for word in cell) for cell in row] for row in group], rectangle


def pdf_assets(content, *, raster=False):
    from app.application.services.raster_document import RasterDocument
    result = []
    started = time.monotonic()
    with fitz.open(stream=content, filetype='pdf') as doc:
        for number, page in enumerate(doc, 1):
            check_budget(result, started)
            if raster:
                embedded = doc.extract_image(page.get_images(full=True)[0][0])
                data = embedded['image']
                kind, _, metadata, _ = RasterDocument.image(data)
                result.append(make_asset(data, content, f'page:{number}:image:0', 'image',
                                         SourceLocation(kind='page', number=number, region=(0, 0, 1, 1)),
                                         'image/' + kind, image_candidates(metadata)))
            else:
                # Existing PDF subset has no drawings/images/interactive surfaces.
                # Text alignment is only a candidate table, requiring human review.
                if page.rotation or page.cropbox != page.mediabox:
                    raise failure()
                for index, (rows, rectangle) in enumerate(aligned_tables(page)):
                    from app.application.services.local_image_screening import LocalImageScreening
                    if page.rect.width * page.rect.height * 4 > LocalImageScreening.MAX_PIXELS:
                        raise failure()
                    try:
                        inspected = LocalImageScreening.inspect(
                            page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes('png')
                        )
                    except WorkflowError as exc:
                        if exc.code != 'SOURCE_SCREENING_INCOMPLETE':
                            raise
                        # Native-text ingestion is already screened separately.
                        # If optional visual qualification is unavailable, omit
                        # this derived table asset instead of admitting an
                        # unreviewable graphic or rejecting the text resource.
                        continue
                    if inspected.faces or not inspected.lines:
                        raise failure()
                    screen_raster_text('\n'.join(line.text for line in inspected.lines))
                    screen_raster_text(page.get_text())
                    screen_raster_text('\n'.join(str(value) for value in (doc.metadata or {}).values()))
                    x0, y0, x1, y1 = rectangle
                    asset = make_asset(table_bytes(rows), content, f'page:{number}:table:{index}', 'table',
                                       SourceLocation(kind='page', number=number, region=(x0 / page.rect.width, y0 / page.rect.height,
                                                                                       x1 / page.rect.width, y1 / page.rect.height)),
                                       'application/json')
                    result.append(unattributed_fields(asset, page.get_text()))
                    check_budget(result, started)
    check_budget(result, started)
    return result


def review_digest(resource):
    return hashlib.sha256(json.dumps([resource.id, resource.metadata.original_sha256,
        [asset.model_dump(mode='json') for asset in resource.metadata.assets],
        [row.model_dump(mode='json') for row in resource.metadata.asset_reviews]], sort_keys=True).encode()).hexdigest()


def is_reviewed(resource):
    return bool(resource.metadata.asset_reviews and resource._asset_review_receipt == review_digest(resource))


def require_reviewed(resources):
    for resource in resources:
        if resource.metadata.assets and not is_reviewed(resource):
            raise WorkflowError('ASSET_CONFIRMATION_REQUIRED', 'Review extracted images and tables against their original before using this source.')


def apply_reviews(original, reviews):
    result = original.model_copy(deep=True)
    assets = {asset.id: asset for asset in result.metadata.assets}
    revisions = sorted({row.revision for row in reviews})
    if not assets or not revisions or revisions != list(range(1, len(revisions) + 1)):
        raise failure()
    for revision in revisions:
        rows = [row for row in reviews if row.revision == revision]
        if sorted(row.asset_id for row in rows) != sorted(assets):
            raise failure()
        for row in rows:
            asset = assets[row.asset_id]
            if (row.original_sha256 != original.metadata.original_sha256 or row.asset_sha256 != asset.sha256
                    or row.initial != {name: getattr(asset, name) for name in FIELDS}
                    or set(row.corrected) != set(FIELDS)):
                raise failure()
            for name, value in row.corrected.items():
                field = asset.fields[name]
                # A correction is an explicit selection among source-attributed
                # values, or withdrawal. No unsupported free-text bibliography.
                if value is not None:
                    screen_raster_text(value)
                    if field.status in ('not_attributable', 'ambiguous') or value not in field.values:
                        raise failure()
                setattr(asset, name, value)
            asset.review_status = 'confirmed'
    result.metadata.asset_reviews = reviews
    result._asset_review_receipt = review_digest(result)
    return result


def asset_content(resource, asset_id, *, original_region=False):
    from app.application.services.source_document import SourceDocument
    SourceDocument.verify(resource, resource._original_content)
    asset = next((item for item in resource.metadata.assets if item.id == asset_id), None)
    if asset is None:
        raise failure()
    if original_region and resource.file_type.value == 'pdf':
        with fitz.open(stream=resource._original_content, filetype='pdf') as doc:
            page = doc[asset.location.number - 1]
            x0, y0, x1, y1 = asset.location.region
            clip = fitz.Rect(x0 * page.rect.width, y0 * page.rect.height, x1 * page.rect.width, y1 * page.rect.height)
            if clip.width * clip.height * 4 > 12_000_000:
                raise failure()
            return page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=clip, alpha=False).tobytes('png'), 'image/png'
    # PPTX: original shape bytes/cells, not a claim of rendered-slide fidelity.
    return base64.b64decode(asset.content_base64, validate=True), asset.media_type
