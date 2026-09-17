"""Restricted OOXML evidence reader. All ZIP/XML/image work stays in memory.

This is not a renderer, template importer, or a general privacy certification.
Opaque parts and unknown XML vocabulary fail closed. Asset descriptors point to
original slide shapes; no reusable image/table files or OCR evidence are made.
"""
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import posixpath
import re
import stat
import struct
import time
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED, ZIP_STORED

from pptx import Presentation

from app.application.services.local_image_screening import LocalImageScreening
from app.application.services.prototype_policy import PrototypePolicy
from app.application.services.raster_document import RasterDocument
from app.application.services.raster_privacy import screen_raster_text
from app.application.services.source_date_policy import SourceDatePolicy
from app.application.services.source_screening import SourceScreening
from app.domain.exceptions.workflow_error import WorkflowError
from app.domain.models.resource import Resource
from app.domain.models.source_metadata import SourceAsset, SourceLocation, SourceMetadata

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
CP = 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'
DC = 'http://purl.org/dc/elements/1.1/'
EP = 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'

# An explicit vocabulary, not "anything in an Office namespace". No animation,
# AlternateContent, custom geometry, charts, embedded objects, links or actions.
VOCABULARY = {
    P: '''presentation sldMasterIdLst sldMasterId sldIdLst sldId sldSz notesSz
        notesMasterIdLst notesMasterId defaultTextStyle sld cSld spTree nvGrpSpPr
        cNvPr cNvGrpSpPr nvPr grpSpPr sp nvSpPr cNvSpPr spPr txBody clrMapOvr
        sldMaster sldLayoutIdLst sldLayoutId txStyles titleStyle bodyStyle otherStyle
        sldLayout ph clrMap bg bgRef bgPr notes notesMaster notesStyle hf
        pic nvPicPr cNvPicPr blipFill graphicFrame nvGraphicFramePr cNvGraphicFramePr
        xfrm style extLst ext presentationPr viewPr normalViewPr restoredLeft
        restoredTop slideViewPr cSldViewPr cViewPr scale sx sy origin guideLst guide
        notesTextViewPr gridSpacing''',
    A: '''sx sy spPr off ext chOff chExt xfrm prstGeom avLst noFill solidFill srgbClr schemeClr
        sysClr prstClr alpha tint shade satMod lumMod lumOff ln prstDash round
        headEnd tailEnd miter bevel bodyPr lstStyle p pPr r rPr t br endParaRPr
        defPPr defRPr latin ea cs buChar buFont buNone buSzPct buSzPts buClr
        lvl1pPr lvl2pPr lvl3pPr lvl4pPr lvl5pPr lvl6pPr lvl7pPr lvl8pPr lvl9pPr
        spcBef spcAft spcPct spcPts lnSpc tabLst tab normAutofit spAutoFit noAutofit
        spLocks picLocks graphicFrameLocks masterClrMapping fld
        graphic graphicData tbl tblPr tblGrid gridCol tr tc txBody tcPr tableStyleId
        lnL lnR lnT lnB lnTlToBr lnBlToTr cell3D
        blip stretch fillRect srcRect theme themeElements clrScheme dk1 lt1 dk2 lt2
        accent1 accent2 accent3 accent4 accent5 accent6 hlink folHlink fontScheme
        majorFont minorFont font fmtScheme fillStyleLst lnStyleLst effectStyleLst
        effectStyle effectLst outerShdw scene3d camera lightRig rot sp3d bevelT
        bgFillStyleLst gradFill gsLst gs lin path fillToRect objectDefaults spDef
        lnDef style lnRef fillRef effectRef fontRef extraClrSchemeLst tblStyleLst''',
    CP: 'coreProperties category contentStatus creator description keywords lastModifiedBy revision subject title version',
    DC: 'title subject creator description language identifier',
    'http://purl.org/dc/terms/': 'created modified',
    EP: '''Properties TotalTime Words Application PresentationFormat Paragraphs Slides
        Notes HiddenSlides MMClips ScaleCrop HeadingPairs TitlesOfParts Manager Company
        LinksUpToDate SharedDoc HyperlinkBase HyperlinksChanged AppVersion''',
    'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes': 'vector variant lpstr i4',
    'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties': 'Properties property',
    'http://schemas.microsoft.com/office/powerpoint/2010/main': 'creationId defaultImageDpi discardImageEditData',
    REL: 'Relationships Relationship',
    CT: 'Types Default Override',
}
ALLOWED_ATTRIBUTES = set("""ContentType Extension Id PartName Target TargetMode Type
    accent1 accent2 accent3 accent4 accent5 accent6 algn anchor ang autoCompressPictures
    b bIns bandRow baseType bg1 bg2 blurRad cap char cmpd cx cy d def defTabSz dir dist
    eaLnBrk firstRow folHlink h hangingPunct hlink id idx indent kern l lIns lang lastClr
    lastView lat latinLnBrk lon lvl marL n name noChangeAspect noGrp noRot orient path
    pos preserve prst r rIns rev rig rotWithShape rtl rtlCol saveSubsetFonts scaled
    script size smtClean snapToGrid snapToObjects sz t tIns tx1 tx2 txBox type typeface
    uri val varScale vert w wrap x y descr title fmtid pid cstate bwMode flipH flipV
    rot show showMasterSp showMasterPhAnim hidden dirty err kumimoji normalizeH baseline
    gridSpan rowSpan hMerge vMerge bandCol firstCol lastCol lastRow marR marT marB
    fontScale lnSpcReduction useBgFill anchorCtr numCol spcCol compatLnSpc""".split())
ALLOWED_TAGS = {f'{{{ns}}}{name}' for ns, names in VOCABULARY.items() for name in names.split()}
REL_TYPES = {R + '/' + kind for kind in (
    'officeDocument', 'extended-properties', 'custom-properties', 'slide', 'slideMaster',
    'slideLayout', 'notesSlide', 'notesMaster', 'theme', 'presProps', 'viewProps', 'tableStyles', 'image',
)} | {REL + '/metadata/core-properties', REL + '/metadata/thumbnail'}
PARTS = {
    '[Content_Types].xml': (f'{{{CT}}}Types', None),
    'docProps/core.xml': (f'{{{CP}}}coreProperties', 'application/vnd.openxmlformats-package.core-properties+xml'),
    'docProps/app.xml': (f'{{{EP}}}Properties', 'application/vnd.openxmlformats-officedocument.extended-properties+xml'),
    'docProps/custom.xml': ('{http://schemas.openxmlformats.org/officeDocument/2006/custom-properties}Properties', 'application/vnd.openxmlformats-officedocument.custom-properties+xml'),
    'ppt/presentation.xml': (f'{{{P}}}presentation', MIME + '.main+xml'),
    'ppt/presProps.xml': (f'{{{P}}}presentationPr', 'application/vnd.openxmlformats-officedocument.presentationml.presProps+xml'),
    'ppt/viewProps.xml': (f'{{{P}}}viewPr', 'application/vnd.openxmlformats-officedocument.presentationml.viewProps+xml'),
    'ppt/tableStyles.xml': (f'{{{A}}}tblStyleLst', 'application/vnd.openxmlformats-officedocument.presentationml.tableStyles+xml'),
}
for folder, stem, root in (('slides', 'slide', 'sld'), ('slideLayouts', 'slideLayout', 'sldLayout'),
                           ('slideMasters', 'slideMaster', 'sldMaster'), ('notesSlides', 'notesSlide', 'notes'),
                           ('notesMasters', 'notesMaster', 'notesMaster'), ('theme', 'theme', 'theme')):
    PARTS[f'ppt/{folder}/{stem}{{number}}.xml'] = (
        f'{{{A if stem == "theme" else P}}}{root}',
        'application/vnd.openxmlformats-officedocument.' + ('theme' if stem == 'theme' else 'presentationml.' + stem) + '+xml',
    )


class PptxDocument:
    MAX_ENTRIES = 2000
    MAX_EXPANDED_BYTES = 40 * 1024 * 1024
    MAX_PART_BYTES = 8 * 1024 * 1024
    MAX_RATIO = 200
    MAX_XML_NODES = 100000
    MAX_SLIDES = 200
    MAX_IMAGES = 20
    MAX_SECONDS = 60

    @staticmethod
    def incomplete():
        return WorkflowError('SOURCE_SCREENING_INCOMPLETE',
                             'PPTX screening could not complete. Use a dated presentation with supported text, simple tables and inspectable text-only images; remove opaque or interactive content.')

    @classmethod
    def _part(cls, name):
        if name in PARTS:
            return PARTS[name]
        for pattern, descriptor in PARTS.items():
            if '{number}' in pattern and re.fullmatch(re.escape(pattern).replace(r'\{number\}', '[1-9][0-9]*'), name):
                return descriptor
        if name == '_rels/.rels' or re.fullmatch(r'ppt/(?:[A-Za-z]+/)?_rels/[A-Za-z]+[0-9]*\.xml\.rels', name):
            return f'{{{REL}}}Relationships', 'application/vnd.openxmlformats-package.relationships+xml'
        if re.fullmatch(r'(?:ppt/media/[A-Za-z0-9_-]+|docProps/thumbnail)\.(?:png|jpeg|jpg)', name):
            return None, 'image/png' if name.endswith('.png') else 'image/jpeg'
        raise cls.incomplete()

    @classmethod
    def _xml(cls, content, expected_root):
        text = content.decode('utf-8', errors='strict')
        without_declaration = re.sub(r'^\ufeff?<\?xml[^?]*\?>', '', text)
        if '<!' in without_declaration or '<?' in without_declaration:
            raise cls.incomplete()
        parser = ET.XMLParser(target=ET.TreeBuilder(insert_comments=True, insert_pis=True))
        root = ET.fromstring(text, parser=parser)
        if root.tag != expected_root:
            raise cls.incomplete()
        nodes = list(root.iter())
        if len(nodes) > cls.MAX_XML_NODES:
            raise cls.incomplete()
        for node in nodes:
            if node.tag not in ALLOWED_TAGS:
                raise cls.incomplete()
            PrototypePolicy.screen([node.text, node.tail, list(node.attrib.values())])
            # Only exact standard labels in application properties are exempt
            # from the additional possible-full-name heuristic, never markers.
            standard_labels = {'Microsoft Macintosh PowerPoint', 'Office Theme', 'Slide Titles'}
            if node.text and node.text.strip() not in standard_labels:
                screen_raster_text(node.text)
            if node.tail:
                screen_raster_text(node.tail)
            for key, value in node.attrib.items():
                if not key.startswith('{') and key not in ALLOWED_ATTRIBUTES:
                    raise cls.incomplete()
                if key.startswith('{') and not key.startswith(('{' + R + '}', '{http://www.w3.org/2001/XMLSchema-instance}')):
                    raise cls.incomplete()
                if key in ('descr', 'title', 'name'):
                    # Default placeholder names are technical labels. Other shape
                    # names, alternate text and custom property names are screened.
                    standard_name = (key == 'name' and (
                        value in {'Office Theme', 'Office', 'Blank', 'Title Slide', 'Title and Content',
                                  'Section Header', 'Two Content', 'Comparison', 'Title Only',
                                  'Content with Caption', 'Picture with Caption',
                                  'Title and Vertical Text', 'Vertical Title and Text', 'Vertical Title 1'}
                        or re.fullmatch(r'(?:Title|Text|Vertical Text|Picture|Content|Date|Footer|Slide Number|Slide Image|Notes|Header) Placeholder [0-9]+', value)))
                    if not standard_name:
                        screen_raster_text(value)
            # Reassemble split runs so XML markup cannot split an identifier.
            if node.tag == f'{{{A}}}p':
                screen_raster_text(''.join(n.text or '' for n in node.iter(f'{{{A}}}t')))
            if node.tag in (f'{{{DC}}}creator', f'{{{CP}}}lastModifiedBy', f'{{{EP}}}Manager', f'{{{EP}}}Company'):
                screen_raster_text(' '.join(node.itertext()))
            if node.tag == f'{{{A}}}prstGeom' and node.get('prst') != 'rect':
                raise cls.incomplete()
            if node.tag == f'{{{A}}}graphicData' and node.get('uri') != 'http://schemas.openxmlformats.org/drawingml/2006/table':
                raise cls.incomplete()
            if node.tag in (f'{{{P}}}pic', f'{{{P}}}graphicFrame') and expected_root != f'{{{P}}}sld':
                raise cls.incomplete()
            if node.tag == f'{{{P}}}sp':
                properties = node.find(f'{{{P}}}spPr')
                if properties is not None and any(
                        child.tag not in {f'{{{A}}}{tag}' for tag in ('xfrm', 'prstGeom', 'noFill')}
                        and not (child.tag == f'{{{A}}}ln' and len(child) == 1 and (child[0].tag == f'{{{A}}}noFill'
                            or (expected_root == f'{{{P}}}notesMaster'
                                and node.find(f'.//{{{P}}}ph[@type="sldImg"]') is not None
                                and child.find(f'{{{A}}}solidFill/{{{A}}}prstClr[@val="black"]') is not None)))
                        for child in properties):
                    raise cls.incomplete()
            if node.tag in (f'{{{P}}}sp', f'{{{P}}}pic'):
                # Text boxes/placeholders and pictures only; no composed vector art.
                if node.tag == f'{{{P}}}sp' and node.find(f'.//{{{P}}}ph') is None and node.find(f'.//{{{P}}}cNvSpPr[@txBox="1"]') is None:
                    raise cls.incomplete()
        cls._structure(root)
        PrototypePolicy.screen('\n'.join(root.itertext()))
        if expected_root in (f'{{{P}}}sld', f'{{{P}}}notes'):
            screen_raster_text('\n'.join(''.join(p.itertext()) for p in root.iter(f'{{{A}}}p')))
        return root

    @classmethod
    def _structure(cls, root):
        """Validate cardinalities that otherwise produce silent first-match reads."""
        required = {
            f'{{{P}}}presentation': [f'{{{P}}}{tag}' for tag in ('sldIdLst', 'sldMasterIdLst', 'sldSz')],
            f'{{{P}}}cSld': [f'{{{P}}}spTree'],
            f'{{{P}}}spTree': [f'{{{P}}}nvGrpSpPr', f'{{{P}}}grpSpPr'],
            f'{{{P}}}sp': [f'{{{P}}}{tag}' for tag in ('nvSpPr', 'spPr')],
            f'{{{P}}}pic': [f'{{{P}}}{tag}' for tag in ('nvPicPr', 'blipFill', 'spPr')],
            f'{{{P}}}graphicFrame': [f'{{{P}}}nvGraphicFramePr', f'{{{P}}}xfrm', f'{{{A}}}graphic'],
            f'{{{A}}}graphic': [f'{{{A}}}graphicData'],
            f'{{{A}}}graphicData': [f'{{{A}}}tbl'],
            f'{{{A}}}tbl': [f'{{{A}}}tblPr', f'{{{A}}}tblGrid'],
            f'{{{A}}}tc': [f'{{{A}}}txBody', f'{{{A}}}tcPr'],
        }
        for tag in ('sld', 'sldLayout', 'sldMaster', 'notes', 'notesMaster'):
            required[f'{{{P}}}{tag}'] = [f'{{{P}}}cSld']
        for node in root.iter():
            for tag in required.get(node.tag, []):
                if len(node.findall(tag)) != 1:
                    raise cls.incomplete()
            if node.tag == f'{{{P}}}spTree':
                allowed = {f'{{{P}}}{tag}' for tag in ('nvGrpSpPr', 'grpSpPr', 'sp', 'pic', 'graphicFrame')}
                if any(child.tag not in allowed for child in node):
                    raise cls.incomplete()
            if node.tag == f'{{{P}}}sp' and len(node.findall(f'{{{P}}}txBody')) > 1:
                raise cls.incomplete()
            if node.tag == f'{{{A}}}tbl':
                columns = node.find(f'{{{A}}}tblGrid')
                rows = node.findall(f'{{{A}}}tr')
                if not len(columns) or not rows or any(len(row.findall(f'{{{A}}}tc')) != len(columns) for row in rows):
                    raise cls.incomplete()
            if node.tag == f'{{{P}}}sp' and root.tag in (f'{{{P}}}sldMaster', f'{{{P}}}sldLayout', f'{{{P}}}notesMaster'):
                if node.find(f'.//{{{P}}}ph') is None:
                    raise cls.incomplete()
            if node.tag == f'{{{A}}}xfrm' and (node.get('rot', '0') != '0' or node.get('flipH', '0') != '0' or node.get('flipV', '0') != '0'):
                raise cls.incomplete()
        shape_ids = [node.get('id') for node in root.iter(f'{{{P}}}cNvPr')]
        if len(shape_ids) != len(set(shape_ids)) or None in shape_ids:
            raise cls.incomplete()

    @classmethod
    def _package(cls, content, start):
        data, roots, types = {}, {}, {}
        # Classic single-disk ZIP only. Refuse preambles, trailers, ZIP64,
        # uninspected local extras and gaps between local records.
        if not content.startswith(b'PK\x03\x04') or content[-22:-18] != b'PK\x05\x06':
            raise cls.incomplete()
        _, disk, directory_disk, disk_count, count, size, offset, comment = struct.unpack('<4s4H2LH', content[-22:])
        if disk or directory_disk or disk_count != count or comment or offset + size != len(content) - 22:
            raise cls.incomplete()
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if archive.comment or not 0 < len(entries) <= cls.MAX_ENTRIES:
                raise cls.incomplete()
            if count != len(entries):
                raise cls.incomplete()
            cursor = 0
            for entry in sorted(entries, key=lambda item: item.header_offset):
                if entry.header_offset != cursor or content[cursor:cursor + 4] != b'PK\x03\x04':
                    raise cls.incomplete()
                flags = struct.unpack_from('<H', content, cursor + 6)[0]
                name_size, extra_size = struct.unpack_from('<HH', content, cursor + 26)
                if flags != entry.flag_bits or flags & ~0x800 or extra_size:
                    raise cls.incomplete()
                cursor += 30 + name_size + entry.compress_size
            if cursor != offset:
                raise cls.incomplete()
            names = [entry.filename for entry in entries]
            if len({name.casefold() for name in names}) != len(names):
                raise cls.incomplete()
            if sum(entry.file_size for entry in entries) > cls.MAX_EXPANDED_BYTES:
                raise cls.incomplete()
            for entry in entries:
                name = entry.filename
                if (entry.orig_filename != name or entry.flag_bits & 1 or entry.extra or entry.comment
                        or entry.compress_type not in (ZIP_STORED, ZIP_DEFLATED)
                        or stat.S_IFMT(entry.external_attr >> 16) not in (0, stat.S_IFREG)
                        or entry.file_size > cls.MAX_PART_BYTES
                        or entry.file_size > max(1, entry.compress_size) * cls.MAX_RATIO
                        or time.monotonic() - start > cls.MAX_SECONDS):
                    raise cls.incomplete()
                if name != '[Content_Types].xml':
                    screen_raster_text(re.sub(r'[_/.-]+', ' ', name))
                root_tag, types[name] = cls._part(name)
                data[name] = archive.read(entry)  # bounded above; CRC/local-header checks
                if root_tag:
                    roots[name] = cls._xml(data[name], root_tag)
        if sum(sum(1 for _ in root.iter()) for root in roots.values()) > cls.MAX_XML_NODES:
            raise cls.incomplete()
        if not {'[Content_Types].xml', '_rels/.rels', 'ppt/presentation.xml'} <= roots.keys():
            raise cls.incomplete()
        defaults, overrides = {}, {}
        for node in roots['[Content_Types].xml']:
            mapping, key = (defaults, node.get('Extension')) if node.tag == f'{{{CT}}}Default' else (overrides, node.get('PartName'))
            if not key or key in mapping:
                raise cls.incomplete()
            mapping[key] = node.get('ContentType')
        for name, expected in types.items():
            if expected and overrides.get('/' + name, defaults.get(name.rsplit('.', 1)[-1])) != expected:
                raise cls.incomplete()
        if any(name[1:] not in data for name in overrides):
            raise cls.incomplete()
        relationships = {}
        for name, root in roots.items():
            if not name.endswith('.rels'):
                continue
            owner = '' if name == '_rels/.rels' else name.replace('/_rels/', '/')[:-5]
            if owner and owner not in roots:
                raise cls.incomplete()
            relationships[owner] = {}
            for node in root:
                identifier, target, kind = node.get('Id'), node.get('Target', ''), node.get('Type')
                if (kind not in REL_TYPES or node.get('TargetMode', 'Internal') != 'Internal'
                        or not identifier or identifier in relationships[owner]
                        or not target or re.search(r'[:%?#\\\x00-\x20]', target) or target.startswith('/')):
                    raise cls.incomplete()
                target = posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
                if target not in data or target.endswith('.rels') or target == '[Content_Types].xml':
                    raise cls.incomplete()
                suffix = kind.rsplit('/', 1)[-1]
                expected_types = {
                    'officeDocument': MIME + '.main+xml',
                    'core-properties': PARTS['docProps/core.xml'][1],
                    'extended-properties': PARTS['docProps/app.xml'][1],
                    'custom-properties': PARTS['docProps/custom.xml'][1],
                }
                if suffix in ('image', 'thumbnail'):
                    valid_type = types[target] in ('image/png', 'image/jpeg')
                else:
                    expected_type = expected_types.get(suffix, 'application/vnd.openxmlformats-officedocument.' + ('theme' if suffix == 'theme' else 'presentationml.' + suffix) + '+xml')
                    valid_type = types[target] == expected_type
                if not valid_type:
                    raise cls.incomplete()
                relationships[owner][identifier] = (kind, target)
        for name, root in roots.items():
            kinds = [kind.rsplit('/', 1)[-1] for kind, _ in relationships.get(name, {}).values()]
            required_relations = {
                f'{{{P}}}sld': ('slideLayout',), f'{{{P}}}sldLayout': ('slideMaster',),
                f'{{{P}}}sldMaster': ('theme',), f'{{{P}}}notes': ('slide', 'notesMaster'),
                f'{{{P}}}notesMaster': ('theme',),
            }
            if any(kinds.count(kind) != 1 for kind in required_relations.get(root.tag, ())):
                raise cls.incomplete()
            if root.tag == f'{{{P}}}sld' and kinds.count('notesSlide') > 1:
                raise cls.incomplete()
        # No orphaned parts or relationships: nothing can hide outside the graph.
        seen, pending = set(), ['']
        while pending:
            owner = pending.pop()
            for _, target in relationships.get(owner, {}).values():
                if target not in seen:
                    seen.add(target)
                    pending.append(target)
        if seen != {name for name in data if not name.endswith('.rels') and name != '[Content_Types].xml'}:
            raise cls.incomplete()
        office = [target for kind, target in relationships[''].values() if kind == R + '/officeDocument']
        if office != ['ppt/presentation.xml']:
            raise cls.incomplete()
        for name, root in roots.items():
            for node in root.iter():
                for key, value in node.attrib.items():
                    if key.startswith('{' + R + '}') and value not in relationships.get(name, {}):
                        raise cls.incomplete()
        slide_ids = roots['ppt/presentation.xml'].find(f'{{{P}}}sldIdLst')
        if slide_ids is None or not 0 < len(slide_ids) <= min(cls.MAX_SLIDES, SourceScreening.MAX_PAGES):
            raise cls.incomplete()
        ordered = []
        ids = set()
        for node in slide_ids:
            kind, target = relationships['ppt/presentation.xml'][node.get(f'{{{R}}}id')]
            if kind != R + '/slide' or target in ordered or node.get('id') in ids:
                raise cls.incomplete()
            ids.add(node.get('id'))
            ordered.append(target)
        if set(ordered) != {name for name in roots if re.fullmatch(r'ppt/slides/slide[0-9]+\.xml', name)}:
            raise cls.incomplete()
        images = [name for name, root in types.items() if root in ('image/png', 'image/jpeg')]
        if len(images) > cls.MAX_IMAGES:
            raise cls.incomplete()
        for name in images:
            if time.monotonic() - start > cls.MAX_SECONDS:
                raise cls.incomplete()
            kind, png, _, _ = RasterDocument.image(data[name])
            if types[name] != 'image/' + kind:
                raise cls.incomplete()
            result = LocalImageScreening.inspect(png)
            if result.faces or not result.lines:
                raise cls.incomplete()
            screen_raster_text('\n'.join(line.text for line in result.lines))
        if time.monotonic() - start > cls.MAX_SECONDS:
            raise cls.incomplete()
        return roots, ordered

    @classmethod
    def _date(cls, pages, roots):
        evidence = SourceDatePolicy.derive(pages)
        # Front-slide evidence is required. Dates on other slides/notes/custom
        # properties can contradict it, but cannot rescue an undated front slide.
        lines = [line for page in pages for line in page['text'].splitlines()
                 if SourceDatePolicy._line.fullmatch(line.strip())]
        for root in roots.values():
            for node in root.iter():
                if node.tag == f'{{{A}}}p':
                    value = ''.join(n.text or '' for n in node.iter(f'{{{A}}}t')).strip()
                    if SourceDatePolicy._line.fullmatch(value):
                        lines.append(value)
                elif node.tag.endswith('}property'):
                    name = node.get('name', '').lower().replace(' ', '').replace('_', '')
                    if name in ('publicationdate', 'lastupdated', 'lastupdate', 'updated'):
                        label = 'Published' if name == 'publicationdate' else 'Updated'
                        lines.append(label + ': ' + ''.join(node.itertext()).strip())
        # derive scans only 40 lines by design: explicitly compare each candidate.
        for line in lines:
            candidate = SourceDatePolicy.derive([{'page': 1, 'text': line}])
            combined = [{'page': 1, 'text': evidence.excerpt + '\n' + candidate.excerpt}]
            selected = SourceDatePolicy.derive(combined)
            if selected != evidence:
                raise WorkflowError('SOURCE_DATE_CONFLICT', 'Scientific dates outside the front slide conflict with its date. Replace the source with an unambiguous version.')
        return evidence

    @classmethod
    def read(cls, filename, content, resource_id):
        try:
            if not content or len(content) > SourceScreening.MAX_BYTES or not filename.lower().endswith('.pptx'):
                raise cls.incomplete()
            screen_raster_text(re.sub(r'[_/.-]+', ' ', filename))
            start = time.monotonic()
            roots, ordered = cls._package(content, start)
            presentation = Presentation(BytesIO(content))
            if len(presentation.slides) != len(ordered):
                raise cls.incomplete()
            width, height = presentation.slide_width, presentation.slide_height
            if not width or not height:
                raise cls.incomplete()
            pages, assets = [], []
            for number, slide in enumerate(presentation.slides, 1):
                text = []
                shape_ids = set()
                for shape in slide.shapes:
                    if shape.shape_id in shape_ids:
                        raise cls.incomplete()
                    shape_ids.add(shape.shape_id)
                    if shape.has_text_frame:
                        text.append(shape.text)
                    if shape.has_table or shape._element.tag == f'{{{P}}}pic':
                        box = (shape.left / width, shape.top / height,
                               (shape.left + shape.width) / width, (shape.top + shape.height) / height)
                        if not (0 <= box[0] < box[2] <= 1 and 0 <= box[1] < box[3] <= 1):
                            raise cls.incomplete()
                        kind = 'table' if shape.has_table else 'image'
                        assets.append(SourceAsset(id=f'slide-{number}-shape-{shape.shape_id}', kind=kind,
                                                  media_type=None if shape.has_table else shape.image.content_type,
                                                  location=SourceLocation(kind='slide', number=number, region=box)))
                        if shape.has_table:
                            for row in shape.table.rows:
                                if any(cell.is_spanned or cell.is_merge_origin for cell in row.cells):
                                    raise cls.incomplete()
                                text.append('\t'.join(cell.text for cell in row.cells))
                value = '\n'.join(text).strip()
                if not value:
                    raise cls.incomplete()
                screen_raster_text(value)
                pages.append({'page': number, 'text': value})
            evidence = cls._date(pages, roots)
            # Technical created/modified dates are screened but never scientific dates.
            core = roots.get('docProps/core.xml')
            def property_value(tag):
                node = core.find(tag) if core is not None else None
                return node.text.strip() if node is not None and node.text and node.text.strip() else None
            source = Resource(
                id=resource_id, filename=filename, file_type='pptx',
                title=property_value(f'{{{DC}}}title') or filename,
                source=property_value(f'{{{DC}}}creator'),
                extracted_pages=pages, extracted_text='\n'.join(page['text'] for page in pages),
                is_validated=True, uploaded_at=datetime.now(timezone.utc),
                metadata=SourceMetadata(
                    origin='pptx_memory_import', media_type=MIME,
                    author=property_value(f'{{{DC}}}creator'),
                    locations=[SourceLocation(kind='slide', number=n) for n in range(1, len(pages) + 1)],
                    assets=assets, scientific_date=evidence,
                    original_sha256=hashlib.sha256(content).hexdigest(), original_size=len(content),
                ),
            )
            SourceScreening.resource(source, require_text=True)
            source._original_content = content
            return source
        except WorkflowError:
            raise
        except Exception:
            raise cls.incomplete() from None
