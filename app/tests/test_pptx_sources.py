"""05.1: synthetic OOXML evidence, never a presentation-template import."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
import hashlib

from pptx import Presentation
from pptx.util import Inches
import pytest

from app.application.services.source_document import SourceDocument
from app.application.services.source_date_policy import SourceDatePolicy
from app.interfaces.storage.user_session_repository import UserSessionRepository
from app.tests import test_source_lifecycle as lifecycle

workspace = lifecycle.workspace
MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


def rewrite(content, changes=None, remove=()):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name not in remove:
                target.writestr(name, (changes or {}).get(name, source.read(name)))
        for name, data in (changes or {}).items():
            if name not in source.namelist():
                target.writestr(name, data)
    return output.getvalue()


def deck(date='Publication date: 2024', *, notes='Synthetic teaching notes', image=None):
    presentation = Presentation()
    presentation.core_properties.title = 'Synthetic teaching'
    presentation.core_properties.author = ''
    presentation.core_properties.last_modified_by = ''
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    first.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = 'Synthetic evidence\n' + date
    first.notes_slide.notes_text_frame.text = notes
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = 'Synthetic second-slide evidence'
    table = second.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    for row, values in zip(table.rows, [('group', 'value'), ('synthetic', '25')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    if image:
        second.shapes.add_picture(BytesIO(image), Inches(1), Inches(5), width=Inches(2))
    output = BytesIO()
    presentation.save(output)
    # Construct a valid minimal synthetic package without opaque printer settings
    # or the stock binary thumbnail. Production must refuse these, never strip them.
    changes = {}
    removed = ('ppt/printerSettings/printerSettings1.bin', 'docProps/thumbnail.jpeg')
    with ZipFile(BytesIO(output.getvalue())) as source:
        for name in ('[Content_Types].xml', '_rels/.rels', 'ppt/_rels/presentation.xml.rels'):
            root = ET.fromstring(source.read(name))
            for element in list(root):
                if ('printerSettings' in str(element.attrib) or 'thumbnail' in str(element.attrib)
                        or element.get('Extension') == 'bin'):
                    root.remove(element)
            changes[name] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    return rewrite(output.getvalue(), changes, removed)


def upload(client, content, filename='synthetic.pptx', media_type=MIME, route='/resources/synthetic-owner/demo'):
    return client.post(route, files={'file': (filename, content, media_type)})


def test_native_text_table_and_slide_order_are_evidence_not_a_template():
    content = deck()
    source = SourceDocument.read('synthetic.pptx', content, 'source')
    assert source.file_type.value == 'pptx'
    assert source.metadata.origin == 'pptx_memory_import'
    assert source.metadata.media_type == MIME
    assert source.metadata.original_sha256 == hashlib.sha256(content).hexdigest()
    assert source._original_content == content
    assert [location.kind for location in source.metadata.locations] == ['slide', 'slide']
    assert [page['page'] for page in source.extracted_pages] == [1, 2]
    assert 'Synthetic second-slide evidence' in source.extracted_pages[1]['text']
    assert '25' in source.extracted_pages[1]['text']
    assert 'Synthetic teaching notes' not in source.extracted_text
    assert source.metadata.author is None and source.metadata.rights is None
    asset, = source.metadata.assets
    assert asset.kind == 'table' and asset.location.kind == 'slide' and asset.location.number == 2
    assert asset.location.region is not None
    assert SourceDatePolicy.require(source).value == '2024'
    assert SourceDocument.verify(source, content) == content


def test_pptx_api_restart_original_and_owner_isolation(workspace):
    client, repository = workspace
    content = deck()
    response = upload(client, content)
    assert response.status_code == 200, response.text
    identifier = response.json()['resource_id']
    restored = UserSessionRepository(repository.database_path)
    state = restored.load('synthetic-owner', 'demo')[1]
    source, = state.resource_library
    assert state.presentation is None
    assert source._original_content == content
    assert source.metadata.assets[0].location.number == 2
    assert SourceDocument.verify(source, content) == content
    path = f'/projects/synthetic-owner/demo/resources/{identifier}/original'
    original = client.get(path)
    assert original.content == content and original.headers['content-type'] == MIME
    assert original.headers['cache-control'] == 'no-store'
    token = client.post('/auth/register', json={'user_id': 'other-synthetic', 'password': 'synthetic-password'}).json()['token']
    assert client.get(path, headers={'Authorization': f'Bearer {token}'}).status_code == 403
    assert restored.library_resource('other-synthetic', identifier) is None


def changed_xml(content, part, edit):
    with ZipFile(BytesIO(content)) as archive:
        root = ET.fromstring(archive.read(part))
    edit(root)
    return rewrite(content, {part: ET.tostring(root, encoding='utf-8', xml_declaration=True)})


def attack(kind):
    from app.application.services.pptx_document import A, P, R, REL, DC
    content = deck()
    if kind == 'corrupt':
        return b'PK\x03\x04broken synthetic archive'
    if kind == 'encrypted':
        return b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1EncryptedPackage'
    if kind == 'trailing':
        return content + b'opaque trailing data'
    if kind == 'prefix':
        return b'opaque prefix' + content
    if kind == 'missing_part':
        return rewrite(content, remove=('ppt/slides/slide2.xml',))
    if kind == 'orphan':
        return rewrite(content, {'ppt/slides/slide99.xml': b'<broken/>'})
    if kind == 'zip_traversal':
        return rewrite(content, {'../outside.txt': b'synthetic'})
    if kind == 'zip_bomb':
        return rewrite(content, {'ppt/media/image1.png': b'x' * 1000000})
    if kind == 'duplicate':
        output = BytesIO(content)
        with pytest.warns(UserWarning, match='Duplicate name'):
            with ZipFile(output, 'a') as archive:
                archive.writestr('ppt/slides/slide1.xml', b'<broken/>')
        return output.getvalue()
    if kind == 'binary':
        return rewrite(content, {'ppt/embeddings/oleObject1.bin': b'opaque'})
    if kind == 'dtd':
        with ZipFile(BytesIO(content)) as archive:
            value = archive.read('ppt/slides/slide1.xml').replace(b'<p:sld ', b'<!DOCTYPE a [<!ENTITY x "synthetic">]><p:sld ', 1)
        return rewrite(content, {'ppt/slides/slide1.xml': value})
    if kind in ('metadata_email', 'metadata_name'):
        return changed_xml(content, 'docProps/core.xml', lambda root: setattr(root.find(f'{{{DC}}}description'), 'text',
                           'alice@example.invalid' if kind == 'metadata_email' else 'Alice Example'))
    if kind in ('notes_name', 'notes_email'):
        return deck(notes='Alice Example' if kind == 'notes_name' else 'alice@example.invalid')
    if kind == 'shape_name':
        return changed_xml(content, 'ppt/slides/slide1.xml', lambda root: root.find(f'.//{{{P}}}cNvPr').set('name', 'Alice Example'))
    if kind == 'unknown_attribute':
        return changed_xml(content, 'ppt/slides/slide1.xml', lambda root: root.set('opaquePayload', 'YWxpY2VAZXhhbXBsZS5pbnZhbGlk'))
    if kind == 'unknown_element':
        return changed_xml(content, 'ppt/slides/slide1.xml', lambda root: ET.SubElement(root, f'{{{P}}}unknown'))
    if kind == 'notes_split_contact':
        def split(root):
            paragraph = root.find(f'.//{{{A}}}p')
            paragraph.clear()
            for value in ('alice@', 'example.invalid'):
                ET.SubElement(ET.SubElement(paragraph, f'{{{A}}}r'), f'{{{A}}}t').text = value
        return changed_xml(content, 'ppt/notesSlides/notesSlide1.xml', split)
    if kind in ('external', 'dangling', 'unknown_relation', 'wrong_relation_type'):
        def rel(root):
            node = root[0]
            if kind == 'external':
                node.set('Target', 'https://example.invalid/source')
                node.set('TargetMode', 'External')
            elif kind == 'dangling':
                node.set('Target', 'missing.xml')
            else:
                node.set('Type', R + ('/unknown' if kind == 'unknown_relation' else '/image'))
        return changed_xml(content, 'ppt/slides/_rels/slide1.xml.rels', rel)
    if kind == 'duplicate_slide':
        def duplicate(root):
            items = root.find(f'{{{P}}}sldIdLst')
            items[1].set(f'{{{R}}}id', items[0].get(f'{{{R}}}id'))
        return changed_xml(content, 'ppt/presentation.xml', duplicate)
    if kind == 'date_other_slide':
        return changed_xml(content, 'ppt/slides/slide2.xml', lambda root: setattr(root.find(f'.//{{{A}}}t'), 'text', 'Publication date: 2025'))
    if kind == 'date_notes':
        return deck(notes='Publication date: 2025')
    if kind == 'date_missing':
        return deck(date='No scientific date')
    if kind == 'date_conflict':
        return deck(date='Published: 2024\nPublished: 2025')
    if kind in ('custom_name', 'custom_date'):
        from app.application.services.pptx_document import CT
        custom = b'''<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="description"><vt:lpstr>Alice Example</vt:lpstr></property></Properties>'''
        if kind == 'custom_date':
            custom = custom.replace(b'description', b'publicationDate').replace(b'Alice Example', b'2025')
        content = rewrite(content, {'docProps/custom.xml': custom})
        content = changed_xml(content, '[Content_Types].xml', lambda root: ET.SubElement(root, f'{{{CT}}}Override',
                              PartName='/docProps/custom.xml', ContentType='application/vnd.openxmlformats-officedocument.custom-properties+xml'))
        return changed_xml(content, '_rels/.rels', lambda root: ET.SubElement(root, f'{{{REL}}}Relationship',
                           Id='rId99', Type=R + '/custom-properties', Target='docProps/custom.xml'))
    raise AssertionError(kind)


@pytest.mark.parametrize('kind', [
    'corrupt', 'encrypted', 'trailing', 'prefix', 'missing_part', 'orphan', 'zip_traversal', 'zip_bomb',
    'duplicate', 'binary', 'dtd', 'metadata_email', 'metadata_name', 'notes_name', 'notes_email',
    'shape_name', 'unknown_attribute', 'unknown_element', 'notes_split_contact',
    'external', 'dangling', 'unknown_relation', 'wrong_relation_type', 'duplicate_slide',
    'date_missing', 'date_conflict', 'date_other_slide', 'date_notes', 'custom_name', 'custom_date',
])
def test_refusal_is_atomic_and_diagnostics_contain_no_raw_content(workspace, caplog, kind):
    client, repository = workspace
    before = lifecycle.snapshot(repository)
    files = set(repository.database_path.parent.rglob('*'))
    response = upload(client, attack(kind))
    assert response.status_code == 422, (kind, response.text)
    assert lifecycle.snapshot(repository) == before
    assert set(repository.database_path.parent.rglob('*')) == files
    assert 'Alice Example' not in response.text + caplog.text
    assert 'alice@example.invalid' not in response.text + caplog.text
    assert 'opaquePayload' not in response.text + caplog.text


@pytest.mark.parametrize('filename,content', [
    ('disguised.pptx', lifecycle.pdf()), ('disguised.pdf', deck()),
    ('unsupported.docx', deck()), ('unsupported.xlsx', deck()), ('unsupported.dcm', deck()),
])
def test_direct_format_disguise_is_refused(filename, content):
    from app.domain.exceptions.workflow_error import WorkflowError
    from app.application.use_cases.extract_pdf_resource import ExtractPdfResourceUseCase
    with pytest.raises(WorkflowError):
        ExtractPdfResourceUseCase().execute(filename, content, prototype_declaration='synthetic')


@pytest.mark.parametrize('entry', ['add_project', 'add_presentation', 'save', 'save_event', 'sync', 'remember', 'model'])
@pytest.mark.parametrize('tamper', ['slide', 'asset', 'original_missing'])
def test_direct_boundaries_rederive_from_original_before_mutation(tmp_path, entry, tamper):
    from app.ai.workflows.graph_state import GraphState
    from app.ai.prompt_builders.evidence_context_builder import EvidenceContextBuilder
    from app.application.use_cases.manage_project_resources import AddProjectResourceUseCase
    from app.application.use_cases.add_resource import AddResourceUseCase
    from app.domain.exceptions.workflow_error import WorkflowError
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    if tamper == 'slide':
        source.extracted_pages[1]['text'] = 'Forged scientific passage'
    elif tamper == 'asset':
        source.metadata.assets[0].location.number = 1
    else:
        source._original_content = None
    repository = UserSessionRepository(tmp_path / 'sources.sqlite3')
    state = GraphState(prototype_declaration='synthetic')
    presentation = lifecycle.presentation([])
    before = lifecycle.snapshot(repository)
    with pytest.raises(WorkflowError):
        if entry == 'add_project':
            AddProjectResourceUseCase().execute(state, source)
        elif entry == 'add_presentation':
            AddResourceUseCase().execute(presentation, source)
        elif entry == 'model':
            EvidenceContextBuilder().for_overview([source])
        elif entry == 'sync':
            with repository._connect() as connection:
                repository._sync_resources(connection, 'owner', 'project', [source])
        elif entry == 'remember':
            with repository._connect() as connection:
                repository._remember_source(connection, 'owner', source)
        elif entry == 'save':
            repository.save('owner', 'project', 'thread', GraphState(resource_library=[source]))
        else:
            repository.save_with_event('owner', 'project', 'thread', GraphState(resource_library=[source]), 'RESOURCE_UPLOADED', 'owner')
    assert not state.resource_library and not presentation.resources
    assert lifecycle.snapshot(repository) == before


def test_slide_numbers_follow_presentation_order_not_part_filenames():
    from app.application.services.pptx_document import P
    def reverse(root):
        items = root.find(f'{{{P}}}sldIdLst')
        items[:] = list(reversed(list(items)))
    content = changed_xml(deck(), 'ppt/presentation.xml', reverse)
    # The former second slide is now first; it has no scientific date.
    from app.domain.exceptions.workflow_error import WorkflowError
    with pytest.raises(WorkflowError) as error:
        SourceDocument.read('synthetic.pptx', content, 'source')
    assert error.value.code == 'SOURCE_DATE_REQUIRED'


def test_pptx_permanent_delete_removes_original_inventory_text_chunks_and_fences_stale_saves(workspace):
    from app.domain.exceptions.workflow_error import WorkflowError
    client, repository = workspace
    identifier = upload(client, deck()).json()['resource_id']
    thread, stale = repository.load('synthetic-owner', 'demo')
    result = client.delete(f'/users/synthetic-owner/resources/{identifier}')
    assert result.status_code == 200, result.text
    restored = UserSessionRepository(repository.database_path)
    assert restored.library_resource('synthetic-owner', identifier) is None
    assert restored.load('synthetic-owner', 'demo')[1].resource_library == []
    with restored._connect() as connection:
        for table in ('owner_resources', 'project_resources', 'project_resource_pages', 'project_resource_chunks'):
            assert connection.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0] == 0
    raw = lifecycle.snapshot(restored)
    assert 'Synthetic second-slide evidence' not in raw and 'slide-2-shape-' not in raw
    with pytest.raises(WorkflowError):
        restored.save('synthetic-owner', 'demo', thread, stale)


@pytest.mark.parametrize('limit', ['MAX_ENTRIES', 'MAX_EXPANDED_BYTES', 'MAX_PART_BYTES', 'MAX_RATIO', 'MAX_XML_NODES', 'MAX_SLIDES', 'MAX_SECONDS'])
def test_package_limits_fail_before_writes(workspace, monkeypatch, limit):
    from app.application.services.pptx_document import PptxDocument
    monkeypatch.setattr(PptxDocument, limit, 0)
    client, repository = workspace
    before = lifecycle.snapshot(repository)
    assert upload(client, deck()).status_code == 422
    assert lifecycle.snapshot(repository) == before


def monochrome_image():
    from PIL import Image
    output = BytesIO()
    Image.new('RGB', (100, 60), 'white').save(output, format='PNG')
    return output.getvalue()


def test_screened_image_is_inventory_only_not_unconfirmed_ocr_evidence(workspace, monkeypatch):
    from app.application.services.local_image_screening import LocalImageScreening, ImageInspection
    from app.domain.models.source_metadata import OcrRegion
    calls = []
    def inspect(content):
        calls.append(content)
        return ImageInspection(lines=[OcrRegion(box=(0.1, 0.1, 0.9, 0.9), text='Unconfirmed synthetic raster value', confidence=0.5)], faces=0)
    monkeypatch.setattr(LocalImageScreening, 'inspect', inspect)
    client, repository = workspace
    content = deck(image=monochrome_image())
    response = upload(client, content)
    assert response.status_code == 200, response.text
    source = UserSessionRepository(repository.database_path).load('synthetic-owner', 'demo')[1].resource_library[0]
    assert calls
    assert source._original_content == content
    assert [asset.kind for asset in source.metadata.assets] == ['table', 'image']
    image = source.metadata.assets[1]
    assert image.location.number == 2 and image.location.kind == 'slide'
    assert image.media_type == 'image/png' and image.rights is None
    assert not source.metadata.ocr_regions
    assert 'Unconfirmed synthetic raster value' not in lifecycle.snapshot(repository)


def test_image_screening_unavailable_never_accepts_pptx(workspace, monkeypatch):
    from app.application.services.local_image_screening import LocalImageScreening
    def unavailable(content):
        raise RuntimeError('alice@example.invalid raw engine diagnostic')
    monkeypatch.setattr(LocalImageScreening, 'inspect', unavailable)
    client, repository = workspace
    before = lifecycle.snapshot(repository)
    response = upload(client, deck(image=monochrome_image()))
    assert response.status_code == 422
    assert 'alice@example.invalid' not in response.text
    assert lifecycle.snapshot(repository) == before


def test_positive_presentation_selection_preserves_style_and_model_context_labels_slides(tmp_path):
    from app.application.use_cases.add_resource import AddResourceUseCase
    from app.ai.prompt_builders.evidence_context_builder import EvidenceContextBuilder
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    presentation = lifecycle.presentation([])
    presentation.custom_template_id = 'existing-explicit-template'
    from app.ai.workflows.graph_state import GraphState
    from app.domain.exceptions.workflow_error import WorkflowError
    with pytest.raises(WorkflowError, match='Review extracted'):
        AddResourceUseCase().execute(presentation, source)
    repository = UserSessionRepository(tmp_path / 'review.sqlite3')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    repository.review_assets('owner', 'source', 0,
                             [{name: getattr(asset, name) for name in asset.fields} for asset in source.metadata.assets],
                             [asset.id for asset in source.metadata.assets])
    source = repository.library_resource('owner', 'source')
    theme = presentation.theme
    AddResourceUseCase().execute(presentation, source)
    assert presentation.theme == theme
    assert presentation.custom_template_id == 'existing-explicit-template'
    context = EvidenceContextBuilder().for_overview([source])
    assert 'SLIDE: 2' in context
    assert 'Synthetic teaching notes' not in context


def test_pptx_audit_failure_rolls_back_all_source_copies(tmp_path, monkeypatch):
    from app.ai.workflows.graph_state import GraphState
    repository = UserSessionRepository(tmp_path / 'sources.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    before = lifecycle.snapshot(repository)
    def failure(*args):
        raise RuntimeError('Synthetic audit failure')
    monkeypatch.setattr(repository, '_insert_event', failure)
    with pytest.raises(RuntimeError):
        repository.save_with_event('owner', 'project', 'thread', GraphState(resource_library=[source]), 'RESOURCE_UPLOADED', 'owner')
    assert lifecycle.snapshot(UserSessionRepository(repository.database_path)) == before


def test_api_after_restart_corruption_blocks_download_and_model(workspace):
    client, repository = workspace
    identifier = upload(client, deck()).json()['resource_id']
    with repository._connect() as connection:
        connection.execute('UPDATE project_resources SET original_pdf = ?', (deck(date='Published: 2025'),))
    assert client.get(f'/projects/synthetic-owner/demo/resources/{identifier}/original').status_code == 409
    assert client.post('/projects/synthetic-owner/demo/resources/summary').status_code == 409
    assert client.post('/projects/synthetic-owner/demo/resources/discuss', json={'question': 'Summarize this evidence'}).status_code == 409


@pytest.mark.parametrize('route', ['/resources/synthetic-owner/demo', '/resources/pdf/synthetic-owner/demo'])
def test_upload_aliases_use_same_gate_and_minimal_inventory(workspace, route):
    client, repository = workspace
    response = upload(client, deck(), route=route)
    assert response.status_code == 200
    payload = client.get('/projects/synthetic-owner/demo').json()['resource_library'][0]
    assert payload['location_kind'] == 'slide'
    assert payload['asset_inventory'][0]['location']['number'] == 2
    assert 'extracted_pages' not in payload and 'notes' not in payload
    assert upload(client, deck(notes='alice@example.invalid'), route=route).status_code == 422


@pytest.mark.parametrize('filename', ['unsupported.docx', 'unsupported.xlsx', 'unsupported.dcm', 'disguised.pdf', 'Alice_Example.pptx'])
def test_api_filename_controls(workspace, filename):
    client, repository = workspace
    before = lifecycle.snapshot(repository)
    assert upload(client, deck(), filename=filename).status_code in (415, 422)
    assert lifecycle.snapshot(repository) == before


@pytest.mark.parametrize('kind', ['missing_layout_relation', 'duplicate_text_body', 'duplicate_slide_list', 'uneven_table'])
def test_incomplete_or_ambiguous_structure_is_refused(kind):
    from copy import deepcopy
    from app.application.services.pptx_document import A, P
    from app.domain.exceptions.workflow_error import WorkflowError
    content = deck()
    if kind == 'missing_layout_relation':
        content = changed_xml(content, 'ppt/slides/_rels/slide1.xml.rels', lambda root: root.remove(root[0]))
    elif kind == 'duplicate_text_body':
        def duplicate(root):
            shape = root.find(f'.//{{{P}}}sp')
            shape.append(deepcopy(shape.find(f'{{{P}}}txBody')))
        content = changed_xml(content, 'ppt/slides/slide1.xml', duplicate)
    elif kind == 'duplicate_slide_list':
        content = changed_xml(content, 'ppt/presentation.xml', lambda root: root.append(deepcopy(root.find(f'{{{P}}}sldIdLst'))))
    else:
        content = changed_xml(content, 'ppt/slides/slide2.xml', lambda root: root.find(f'.//{{{A}}}tr').remove(root.find(f'.//{{{A}}}tc')))
    with pytest.raises(WorkflowError):
        SourceDocument.read('synthetic.pptx', content, 'source')


def test_selected_pptx_cannot_smuggle_different_inventory(tmp_path):
    from app.ai.workflows.graph_state import GraphState
    from app.application.services.resource_library import resource_selection
    from app.domain.exceptions.workflow_error import WorkflowError
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    selected = resource_selection(source)
    selected.metadata.assets[0].location.number = 1
    repository = UserSessionRepository(tmp_path / 'sources.sqlite3')
    before = lifecycle.snapshot(repository)
    with pytest.raises(WorkflowError):
        repository.save('owner', 'project', 'thread', GraphState(resource_library=[source], presentation=lifecycle.presentation([selected])))
    assert lifecycle.snapshot(repository) == before
