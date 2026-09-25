"""05.3 synthetic role separation and conservative primary-reference warnings."""
import pytest

from app.application.services.citation_presentation import format_citations_for_display
from app.application.services.source_document import SourceDocument
from app.tests.test_pptx_sources import deck
from app.tests import test_source_lifecycle as lifecycle

workspace = lifecycle.workspace

WARNING = 'Source: user-supplied presentation — primary reference unavailable.'


def test_warning_cannot_be_suppressed_by_supplied_bibliographic_text():
    for mention in ('', 'Reference: synthetic article', 'DOI: 10.0000/synthetic', 'References: ambiguous; not attributable'):
        source = SourceDocument.read('synthetic.pptx', deck(date='Publication date: 2024\n' + mention), 'source-1234')
        assert WARNING in format_citations_for_display('[source-1234, p. 2]', [source])


def test_direct_template_save_refuses_unscreened_bytes(workspace):
    _, repository = workspace
    with pytest.raises(ValueError):
        repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', b'opaque', prototype_declaration='synthetic')
    assert repository.list_presentation_templates('synthetic-owner') == []


def setup_presentation(repository):
    thread, state = repository.load('synthetic-owner', 'demo')
    state.presentation = lifecycle.presentation([])
    state.presentation.prototype_declaration = 'synthetic'
    state.presentation.context.target_slide_count = 5
    state.presentation.context.special_instructions = 'Synthetic teaching only'
    state.presentation.context.professional_scope = 'Synthetic teaching'
    state.presentation.context.is_multidisciplinary = False
    from app.application.services.presentation_context_policy import PresentationContextPolicy
    from app.domain.models.professional_scope_declaration import ProfessionalScopeDeclaration
    state.presentation.professional_scope_declaration = ProfessionalScopeDeclaration(
        declared_role='Professor', delivery_purpose='Synthetic teaching', confirmed_within_scope=True,
        actor_user_id='synthetic-owner', is_multidisciplinary=False,
        context_digest=PresentationContextPolicy.digest(state.presentation.context),
    )
    state.presentation.state.blueprint_validated = True
    state.presentation.state.resources_validated = True
    from app.domain.enums.workflow_status import WorkflowStatus
    state.presentation.state.workflow_status = WorkflowStatus.SLIDE_GENERATION
    repository.save('synthetic-owner', 'demo', thread, state)
    return state


def test_explicit_reuse_restart_no_evidence_approval_and_invalidation(workspace):
    from app.tests.test_pptx_sources import upload
    from app.domain.models.slide import Slide
    from app.interfaces.storage.user_session_repository import UserSessionRepository
    client, repository = workspace
    setup_presentation(repository)
    response = upload(client, deck())
    assert response.status_code == 200
    identifier = response.json()['resource_id']
    thread, state = repository.load('synthetic-owner', 'demo')
    assert state.presentation.custom_template_id is None
    before = state.resource_library[0].model_dump()
    response = client.put('/projects/synthetic-owner/demo/theme', json={
        'custom_template_id': identifier, 'expected_revision': state.project_revision,
    })
    assert response.status_code == 200, response.text
    restored = UserSessionRepository(repository.database_path)
    state = restored.load('synthetic-owner', 'demo')[1]
    assert state.presentation.custom_template_id == identifier
    assert state.presentation.resources == []
    assert state.resource_library[0].model_dump() == before
    assert state.presentation.style_selected
    assert state.presentation.state.blueprint_validated
    assert state.presentation.state.resources_validated
    assert not state.resource_library[0].metadata.asset_reviews
    event = restored.list_events('synthetic-owner', 'demo')[0]
    assert event['actor'] == 'synthetic-owner'
    assert event['payload']['template_resource_id'] == identifier
    assert 'Synthetic evidence' not in str(event)
    assert restored.presentation_template_source('other', identifier) is None
    assert not list(repository.templates_path.iterdir())

    state.presentation.slides = [Slide(slide_number=1, title='Synthetic')]
    repository.save('synthetic-owner', 'demo', thread, state)
    with pytest.raises(ValueError, match='before slide generation'):
        restored.select_presentation_style(
            'synthetic-owner', 'demo', state.project_revision,
            template_id=identifier,
        )


def test_hpa_theme_and_colour_are_explicit_durable_and_stage_gated(workspace):
    from app.domain.enums.presentation_theme import PresentationColour, PresentationTheme
    from app.domain.enums.workflow_status import WorkflowStatus
    client, repository = workspace
    state = setup_presentation(repository)
    path = '/projects/synthetic-owner/demo/theme'

    response = client.put(path, json={
        'theme': 'academic', 'colour': 'blue', 'expected_revision': state.project_revision,
    })
    assert response.status_code == 200, response.text
    restored = type(repository)(repository.database_path).load('synthetic-owner', 'demo')[1]
    assert restored.presentation.theme == PresentationTheme.ACADEMIC
    assert restored.presentation.colour == PresentationColour.BLUE
    assert restored.presentation.style_selected
    assert repository.list_events('synthetic-owner', 'demo')[0]['payload']['colour'] == 'blue'

    restored.presentation.state.workflow_status = WorkflowStatus.AWAITING_SLIDE_APPROVAL
    restored.presentation.state.slides_validated = True
    restored.presentation.state.presentation_validated = True
    from app.domain.models.slide import Slide
    restored.presentation.slides = [Slide(slide_number=1, title='Approved', is_validated=True,
                                          approved_by='synthetic-owner')]
    repository.save_with_event(
        'synthetic-owner', 'demo', response.json()['thread_id'], restored,
        'SLIDE_APPROVED', 'synthetic-owner', {'slide_index': 0},
    )
    restored = repository.load('synthetic-owner', 'demo')[1]
    response = client.put(path, json={
        'colour': 'warm', 'expected_revision': restored.project_revision,
    })
    assert response.status_code == 200, response.text
    restarted = type(repository)(repository.database_path)
    changed = restarted.load('synthetic-owner', 'demo')[1].presentation
    assert changed.style_provenance.source == 'hpa_theme'
    assert len(changed.style_provenance.digest) == 64
    assert not changed.slides[0].is_validated
    assert not changed.state.slides_validated
    assert not changed.state.presentation_validated
    event = restarted.list_events('synthetic-owner', 'demo')[0]
    assert event['payload']['style_provenance_digest'] == changed.style_provenance.digest
    assert event['payload']['provenance_verified'] is True


def test_slide_job_requires_explicit_style_without_calling_provider(workspace):
    client, repository = workspace
    setup_presentation(repository)
    response = client.post('/api/v1/projects/synthetic-owner/demo/slides/jobs')
    assert response.status_code == 409
    assert response.json()['detail']['code'] == 'STYLE_SELECTION_REQUIRED'
    assert repository.active_job('synthetic-owner', 'demo') is None


@pytest.mark.parametrize('roundtrip', [False, True])
def test_direct_save_cannot_select_a_graphic_source(workspace, roundtrip):
    from app.ai.workflows.graph_state import GraphState
    _, repository = workspace
    state = setup_presentation(repository)
    template = repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    state.presentation.custom_template_id = template['id']
    if roundtrip:
        state = GraphState.model_validate_json(state.model_dump_json())
    before = lifecycle.snapshot(repository)
    with pytest.raises(ValueError, match='explicit style'):
        repository.save('synthetic-owner', 'demo', 'thread', state)
    assert lifecycle.snapshot(repository) == before


def test_style_concurrency_and_audit_failure_roll_back(workspace, monkeypatch):
    from app.domain.enums.presentation_theme import PresentationTheme
    from app.domain.exceptions.concurrent_modification_error import ConcurrentModificationError
    _, repository = workspace
    state = setup_presentation(repository)
    before = lifecycle.snapshot(repository)
    with pytest.raises(ConcurrentModificationError):
        repository.select_presentation_style('synthetic-owner', 'demo', state.project_revision - 1, theme=PresentationTheme.ACADEMIC)
    def fail(*args, **kwargs):
        raise RuntimeError('synthetic audit failure')
    monkeypatch.setattr(repository, '_insert_event', fail)
    with pytest.raises(RuntimeError):
        repository.select_presentation_style('synthetic-owner', 'demo', state.project_revision, theme=PresentationTheme.ACADEMIC)
    assert lifecycle.snapshot(repository) == before
    with pytest.raises(RuntimeError):
        repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    assert lifecycle.snapshot(repository) == before
    assert not list(repository.templates_path.iterdir())


@pytest.mark.parametrize('permanent', [False, True])
def test_source_removal_clears_graphic_reuse_without_extra_copy(workspace, permanent):
    _, repository = workspace
    state = setup_presentation(repository)
    template = repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    identifier = template['id']
    thread, state = repository.load('synthetic-owner', 'demo')
    state.resource_library.append(repository.library_resource('synthetic-owner', identifier))
    repository.save('synthetic-owner', 'demo', thread, state)
    thread, state = repository.select_presentation_style('synthetic-owner', 'demo', state.project_revision, template_id=identifier)
    if permanent:
        repository.permanently_delete_source('synthetic-owner', identifier, 'synthetic-owner')
        assert repository.library_resource('synthetic-owner', identifier) is None
        assert repository.list_presentation_templates('synthetic-owner') == []
        assert repository.presentation_template_source('synthetic-owner', identifier) is None
    else:
        state.resource_library = []
        repository.save('synthetic-owner', 'demo', thread, state)
        assert repository.library_resource('synthetic-owner', identifier) is not None
    assert repository.load('synthetic-owner', 'demo')[1].presentation.custom_template_id is None


def test_warning_after_restart_api_context_and_selection(workspace):
    from app.tests.test_pptx_sources import upload
    from app.ai.prompt_builders.evidence_context_builder import EvidenceContextBuilder
    from app.application.services.resource_library import resource_selection
    client, repository = workspace
    assert upload(client, deck()).status_code == 200
    state = repository.load('synthetic-owner', 'demo')[1]
    source = state.resource_library[0]
    assert WARNING in format_citations_for_display(f'[{source.id}, p. 2]', [resource_selection(source)])
    response = client.get('/projects/synthetic-owner/demo')
    assert response.status_code == 200
    assert response.json()['resource_library'][0]['primary_reference_warning'] == WARNING
    # Pending extracted tables still block context, even when used graphically.
    with pytest.raises(ValueError, match='Review'):
        EvidenceContextBuilder().for_resources([source], 'synthetic')
    from app.application.services.source_assets import FIELDS
    reviewed = repository.review_assets('synthetic-owner', source.id, 0, [{field: None for field in FIELDS} for asset in source.metadata.assets], [asset.id for asset in source.metadata.assets])
    source = repository.library_resource('synthetic-owner', source.id)
    assert reviewed is not None
    context = EvidenceContextBuilder().for_resources([source], 'synthetic')
    assert WARNING in context
    assert 'SLIDE:' in context
    assert '10.0000' not in context


def test_graphic_only_import_never_enters_model_evidence(workspace):
    from app.application.services.resource_library import resolve_presentation_resources
    _, repository = workspace
    state = setup_presentation(repository)
    template = repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    thread, state = repository.select_presentation_style('synthetic-owner', 'demo', state.project_revision, template_id=template['id'])
    assert state.resource_library == []
    assert resolve_presentation_resources(state) == []
    repository.save('synthetic-owner', 'demo', thread, state)
    assert repository.load('synthetic-owner', 'demo')[1].presentation.custom_template_id == template['id']
    repository.permanently_delete_source('synthetic-owner', template['id'], 'synthetic-owner')
    remaining = repository.load('synthetic-owner', 'demo')[1].presentation
    assert remaining.custom_template_id is None
    assert remaining.state.blueprint_validated
    assert remaining.state.resources_validated


def test_export_and_validation_use_original_bibliography_and_exact_warning(workspace):
    from io import BytesIO
    from pptx import Presentation as PowerPoint
    from app.application.services.source_assets import FIELDS
    from app.application.use_cases.export_powerpoint import ExportPowerPointUseCase
    from app.domain.models.agenda import Agenda
    from app.domain.models.slide import Slide
    _, repository = workspace
    template = repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    source = repository.library_resource('synthetic-owner', template['id'])
    repository.review_assets('synthetic-owner', source.id, 0,
                            [{field: None for field in FIELDS} for asset in source.metadata.assets],
                            [asset.id for asset in source.metadata.assets])
    source = repository.library_resource('synthetic-owner', source.id)
    presentation = lifecycle.presentation([source])
    presentation.state.resources_validated = presentation.state.presentation_validated = True
    presentation.agenda = Agenda(items=['Synthetic'], is_validated=True)
    presentation.slides = [Slide(slide_number=1, title='Synthetic', content='Synthetic second-slide evidence',
        references=['Invented journal'], reference_details=[{
            'resource_id': source.id, 'page': 2, 'evidence_excerpt': 'Synthetic second-slide evidence',
            'title': 'Invented journal', 'source': 'Invented publisher', 'doi': 'invented-doi',
            'year': 1900, 'primary_reference_verified': True, 'source_warning': '',
        }])]
    output = BytesIO()
    ExportPowerPointUseCase().execute(presentation, output)
    text = '\n'.join(shape.text for slide in PowerPoint(output).slides for shape in slide.shapes if shape.has_text_frame)
    assert text.count(WARNING) == 2
    reference_slide = PowerPoint(output).slides[-1]
    warning_box = next(shape for shape in reference_slide.shapes if shape.has_text_frame and shape.text == WARNING)
    identifier_box = next(shape for shape in reference_slide.shapes if shape.has_text_frame and source.id in shape.text)
    assert warning_box.top >= identifier_box.top + identifier_box.height
    assert 'Invented' not in text
    assert presentation.slides[0].reference_details[0]['primary_reference_verified'] is False
    assert 'doi' not in presentation.slides[0].reference_details[0]
    with pytest.raises(ValueError, match='explicit style'):
        ExportPowerPointUseCase().execute(presentation, BytesIO(), BytesIO(deck()))


@pytest.mark.parametrize('kind', ['corrupt', 'orphan', 'zip_traversal', 'missing_part', 'prefix'])
def test_template_api_uses_same_hostile_package_refusals(workspace, kind):
    from app.tests.test_pptx_sources import attack
    client, repository = workspace
    response = client.post('/users/synthetic-owner/templates',
        files={'file': ('synthetic.pptx', attack(kind), 'application/octet-stream')},
        data={'prototype_declaration': 'synthetic', 'external_processing_acknowledged': 'true'})
    assert response.status_code == 409
    assert repository.list_presentation_templates('synthetic-owner') == []
    assert repository.list_library_resources('synthetic-owner') == []
    assert not list(repository.templates_path.iterdir())


def test_reuse_api_permission_revision_and_blueprint_gate(workspace):
    client, repository = workspace
    state = setup_presentation(repository)
    template = repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    payload = {'custom_template_id': template['id'], 'expected_revision': state.project_revision}
    path = '/projects/synthetic-owner/demo/theme'
    assert client.put(path, json=payload, headers={'Authorization': ''}).status_code == 401
    token = client.post('/auth/register', json={'user_id': 'other-synthetic', 'password': 'synthetic-password'}).json()['token']
    assert client.put(path, json=payload, headers={'Authorization': f'Bearer {token}'}).status_code == 403
    assert client.put(path, json={'custom_template_id': template['id']}).status_code == 422
    thread, state = repository.load('synthetic-owner', 'demo')
    state.presentation.state.blueprint_validated = False
    repository.save('synthetic-owner', 'demo', thread, state)
    assert client.put(path, json=payload).json()['detail']['code'] == 'PROJECT_VERSION_CONFLICT'
    payload['expected_revision'] = state.project_revision
    assert client.put(path, json=payload).json()['detail']['code'] == 'BLUEPRINT_APPROVAL_REQUIRED'


def test_graphic_deletion_audit_failure_preserves_original_and_selection(workspace, monkeypatch):
    _, repository = workspace
    state = setup_presentation(repository)
    template = repository.save_presentation_template('synthetic-owner', 'synthetic.pptx', deck(), prototype_declaration='synthetic')
    repository.select_presentation_style('synthetic-owner', 'demo', state.project_revision, template_id=template['id'])
    before = lifecycle.snapshot(repository)
    def fail(*args, **kwargs):
        raise RuntimeError('synthetic audit failure')
    monkeypatch.setattr(repository, '_insert_event', fail)
    with pytest.raises(RuntimeError):
        repository.permanently_delete_source('synthetic-owner', template['id'], 'synthetic-owner')
    assert lifecycle.snapshot(repository) == before


def test_active_job_and_direct_theme_save_cannot_bypass_explicit_selection(workspace):
    from app.domain.enums.presentation_theme import PresentationTheme
    from app.domain.exceptions.project_job_running_error import ProjectJobRunningError
    _, repository = workspace
    state = setup_presentation(repository)
    state.presentation.theme = PresentationTheme.MIDNIGHT
    with pytest.raises(ValueError, match='explicit style'):
        repository.save('synthetic-owner', 'demo', 'thread', state)
    repository.create_job('synthetic-owner', 'demo', 'conversation')
    with pytest.raises(ProjectJobRunningError):
        repository.select_presentation_style('synthetic-owner', 'demo', state.project_revision, theme=PresentationTheme.MIDNIGHT)


def test_supplied_value_visual_layout_is_strict_durable_and_invalidates_approval(workspace):
    from app.domain.models.slide import Slide
    from app.interfaces.storage.user_session_repository import UserSessionRepository
    client, repository = workspace
    state = setup_presentation(repository)
    state.presentation.slides = [Slide(slide_number=1, title='Synthetic', is_validated=True,
                                       approved_by='synthetic-owner')]
    state.presentation.state.slides_validated = True
    state.presentation.state.presentation_validated = True
    repository.save_with_event(
        'synthetic-owner', 'demo', 'thread', state,
        'SLIDE_APPROVED', 'synthetic-owner', {'slide_index': 0},
    )
    state = repository.load('synthetic-owner', 'demo')[1]
    path = '/projects/synthetic-owner/demo/slides/0/visual'
    response = client.put(path, json={
        'expected_revision': state.project_revision,
        'layout': 'text_left_visual_right',
        'visual': {'kind': 'bar_chart', 'categories': ['A', 'B'],
                   'series': {'Supplied': [1.5, 2.0]}},
    })
    assert response.status_code == 200, response.text
    restored = UserSessionRepository(repository.database_path).load('synthetic-owner', 'demo')[1]
    slide = restored.presentation.slides[0]
    assert slide.layout == 'text_left_visual_right'
    assert slide.visual.categories == ['A', 'B']
    assert slide.visual.series == {'Supplied': [1.5, 2.0]}
    assert not slide.is_validated and slide.approved_by is None
    assert slide.visual_provenance.source == 'supplied_values'
    assert len(slide.visual_provenance.digest) == 64
    assert not restored.presentation.state.slides_validated
    assert not restored.presentation.state.presentation_validated
    event = repository.list_events('synthetic-owner', 'demo')[0]
    assert event['event_type'] == 'SLIDE_VISUAL_UPDATED'
    assert event['payload']['visual_kind'] == 'bar_chart'
    assert event['payload']['visual_provenance_digest'] == slide.visual_provenance.digest
    assert event['payload']['provenance_verified'] is True
    assert client.put(path, json={
        'expected_revision': restored.project_revision,
        'layout': 'text_left_visual_right',
        'visual': {'kind': 'line_chart', 'categories': ['A', 'B'], 'series': {'Missing': [1]}},
    }).status_code == 422


def test_visual_server_refuses_stale_revision_unreviewed_image_and_layout_bypass(workspace):
    from app.domain.models.slide import Slide
    from app.domain.exceptions.concurrent_modification_error import ConcurrentModificationError
    from app.domain.exceptions.workflow_error import WorkflowError
    _, repository = workspace
    state = setup_presentation(repository)
    state.presentation.slides = [Slide(slide_number=1, title='Synthetic')]
    repository.save('synthetic-owner', 'demo', 'thread', state)
    state = repository.load('synthetic-owner', 'demo')[1]
    with pytest.raises(ConcurrentModificationError):
        repository.update_slide_visual('synthetic-owner', 'demo', state.project_revision - 1, 0,
                                       layout='text_only', visual=None)
    with pytest.raises(WorkflowError, match='reviewed'):
        from app.domain.models.slide import SlideVisual
        repository.update_slide_visual(
            'synthetic-owner', 'demo', state.project_revision, 0,
            layout='visual_focus',
            visual=SlideVisual(
                kind='image', resource_id='absent-resource', asset_id='unreviewed-image',
            ),
        )
    with pytest.raises(WorkflowError, match='requires a supplied visual'):
        repository.update_slide_visual('synthetic-owner', 'demo', state.project_revision, 0,
                                       layout='visual_focus', visual=None)
    from app.domain.models.slide import SlideVisual
    repository.update_slide_visual(
        'synthetic-owner', 'demo', state.project_revision, 0,
        layout='visual_focus',
        visual=SlideVisual(kind='table', columns=['Supplied'], rows=[['1']]),
    )
    forged = repository.load('synthetic-owner', 'demo')[1]
    forged.presentation.slides[0].visual = SlideVisual(
        kind='table', columns=['Supplied'], rows=[['2']],
    )
    with pytest.raises(WorkflowError, match='visual provenance'):
        repository.save('synthetic-owner', 'demo', 'thread', forged)


def test_pptx_builds_only_the_supplied_table_values():
    from pptx import Presentation as PowerPoint
    from app.application.use_cases.export_powerpoint import ExportPowerPointUseCase, THEMES
    from app.domain.enums.presentation_theme import PresentationTheme
    from app.domain.models.slide import Slide, SlideVisual
    deck = PowerPoint()
    source = Slide(
        slide_number=1, title='Synthetic supplied values', layout='visual_focus',
        visual=SlideVisual(kind='table', columns=['Group', 'Value'], rows=[['A', '12'], ['B', '18']]),
    )
    ExportPowerPointUseCase()._add_content_slide(deck, source, THEMES[PresentationTheme.CLINICAL], [])
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ['Group', 'Value'], ['A', '12'], ['B', '18'],
    ]
