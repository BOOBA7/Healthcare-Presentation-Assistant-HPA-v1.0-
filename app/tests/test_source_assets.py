"""05.2 synthetic retained assets; OCR doubles do not qualify privacy accuracy."""
import hashlib
import base64

import fitz
import pytest

from app.application.services.source_document import SourceDocument
from app.ai.workflows.graph_state import GraphState
from app.interfaces.storage.user_session_repository import UserSessionRepository
from app.tests.test_pptx_sources import deck
from app.tests import test_raster_sources as raster_fixtures
from app.tests.test_raster_sources import synthetic_image, synthetic_scan

engine_double = raster_fixtures.engine_double


def table_pdf():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((40, 40), 'Publication date: 2024')
        for y, row in zip((100, 125, 150), [('group', 'value'), ('synthetic', '25'), ('control', '12')]):
            for x, value in zip((40, 240), row):
                page.insert_text((x, y), value)
        return doc.tobytes()


@pytest.mark.parametrize('name,make', [('synthetic.pptx', lambda: deck(image=synthetic_image())),
                                      ('synthetic.pdf', synthetic_scan), ('table.pdf', table_pdf)])
def test_retained_assets_are_bound_to_original_after_restart(tmp_path, engine_double, name, make):
    content = make()
    source = SourceDocument.read(name, content, 'source')
    assert source.metadata.assets
    for asset in source.metadata.assets:
        assert asset.sha256 == hashlib.sha256(base64.b64decode(asset.content_base64)).hexdigest()
        assert asset.original_sha256 == source.metadata.original_sha256
        assert asset.location.region
        assert asset.review_status == 'pending'
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    restored = UserSessionRepository(repository.database_path).library_resource('owner', 'source')
    assert restored.metadata.assets == source.metadata.assets
    assert SourceDocument.verify(restored, content) == content


def test_asset_review_atomic_authority_and_model_gate(tmp_path):
    from app.ai.prompt_builders.evidence_context_builder import EvidenceContextBuilder
    from app.domain.exceptions.workflow_error import WorkflowError
    from app.domain.models.resource import Resource
    from app.tests.test_source_lifecycle import snapshot
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    with pytest.raises(WorkflowError, match='Review extracted'):
        EvidenceContextBuilder().for_overview([source])
    values = [{name: getattr(asset, name) for name in asset.fields} for asset in source.metadata.assets]
    repository.review_assets('owner', 'source', 0, values, [a.id for a in source.metadata.assets])
    restored = UserSessionRepository(repository.database_path).library_resource('owner', 'source')
    assert restored.metadata.asset_reviews[0].actor == 'owner'
    assert restored.metadata.asset_reviews[0].confirmed_at.tzinfo
    assert restored.metadata.assets[0].review_status == 'confirmed'
    assert 'synthetic' in EvidenceContextBuilder().for_overview([restored]).lower()
    forged = Resource.model_validate_json(restored.model_dump_json())
    forged._original_content = restored._original_content
    with pytest.raises(WorkflowError):
        EvidenceContextBuilder().for_overview([forged])
    before = snapshot(repository)
    with pytest.raises(WorkflowError):
        repository.review_assets('owner', 'source', 0, values, [a.id for a in source.metadata.assets])
    assert snapshot(repository) == before

from app.tests import test_source_lifecycle as lifecycle  # noqa: E402
workspace = lifecycle.workspace


def test_asset_api_review_preview_and_no_raw_errors(workspace):
    client, repository = workspace
    response = client.post('/resources/synthetic-owner/demo', files={'file': ('synthetic.pptx', deck())})
    identifier = response.json()['resource_id']
    path = f'/users/synthetic-owner/resources/{identifier}'
    form = client.get(path + '/asset-review')
    assert form.status_code == 200
    data = form.json()
    assert form.headers['cache-control'] == 'no-store'
    asset = data['assets'][0]
    preview = client.get(path + '/assets/' + asset['id'])
    assert preview.json()[1] == ['synthetic', '25']
    request = {'expected_revision': 0, 'values': [asset['values']], 'confirmed_assets': [asset['id']]}
    assert client.post(path + '/asset-review', json=request).status_code == 200
    assert client.post(path + '/asset-review', json=request).status_code == 409
    error = client.post(path + '/asset-review', json={**request, 'actor': 'alice@example.invalid'})
    assert error.status_code == 422 and 'alice@' not in error.text
    assert client.get(path.replace('synthetic-owner', 'other') + '/asset-review').status_code == 403


@pytest.mark.parametrize('fault', ['bytes', 'hash', 'coordinates', 'bounds', 'date', 'location', 'direct_review'])
def test_substitution_and_direct_repository_bypasses(tmp_path, fault):
    from app.domain.exceptions.workflow_error import WorkflowError
    from app.tests.test_source_lifecycle import snapshot
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    before = snapshot(repository)
    thread, state = repository.load('owner', 'demo')
    source = state.resource_library[0]
    asset = source.metadata.assets[0]
    if fault == 'bytes':
        asset.content_base64 = base64.b64encode(b'[["forged"]]').decode()
    elif fault == 'hash':
        asset.sha256 = '0' * 64
    elif fault == 'coordinates':
        asset.location.region = (0, 0, 0.1, 0.1)
    elif fault == 'bounds':
        asset.location.region = (-1, 0, 2, 1)
    elif fault == 'date':
        source.metadata.scientific_date.value = '2023'
    elif fault == 'location':
        asset.location.number = 1
    else:
        asset.review_status = 'confirmed'
    with pytest.raises(WorkflowError):
        repository.save('owner', 'demo', thread, state)
    assert snapshot(repository) == before


def test_review_audit_failure_interruption_and_invalidation(tmp_path, monkeypatch):
    from app.tests.test_source_lifecycle import snapshot
    from app.domain.models.resource_analysis import ResourceAnalysis
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    for project in ('one', 'two'):
        repository.save('owner', project, 'thread', GraphState(resource_library=[source],
                        resource_analysis=ResourceAnalysis(summary='Stale synthetic summary', resource_ids=['source'])))
    values = [{name: getattr(asset, name) for name in asset.fields} for asset in source.metadata.assets]
    ids = [asset.id for asset in source.metadata.assets]
    original = repository._insert_event
    before = snapshot(repository)
    for error in (RuntimeError, KeyboardInterrupt):
        def fail(*args):
            raise error('Synthetic interruption')
        monkeypatch.setattr(repository, '_insert_event', fail)
        with pytest.raises(error):
            repository.review_assets('owner', 'source', 0, values, ids)
        assert snapshot(repository) == before
    monkeypatch.setattr(repository, '_insert_event', original)
    repository.review_assets('owner', 'source', 0, values, ids)
    for project in ('one', 'two'):
        assert repository.load('owner', project)[1].resource_analysis is None
    repository.permanently_delete_source('owner', 'source', 'owner')
    assert source.metadata.assets[0].content_base64 not in snapshot(repository)
    assert UserSessionRepository(repository.database_path).library_resource('owner', 'source') is None


def test_actual_concurrent_review_has_one_winner(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from app.domain.exceptions.workflow_error import WorkflowError
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    values = [{name: getattr(asset, name) for name in asset.fields} for asset in source.metadata.assets]
    def submit(_):
        try:
            repository.review_assets('owner', 'source', 0, values, [a.id for a in source.metadata.assets])
            return 'accepted'
        except WorkflowError as error:
            return error.code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(submit, range(2))) == ['ASSET_REVIEW_CONFLICT', 'accepted']


def deck_with_caption(*, description=None, title=None, nearby=None, duplicate=False):
    from app.tests.test_pptx_sources import changed_xml
    from app.application.services.pptx_document import P, A
    from copy import deepcopy
    def edit(root):
        shape = root.find(f'.//{{{P}}}graphicFrame')
        properties = shape.find(f'.//{{{P}}}cNvPr')
        for key, value in [('descr', description), ('title', title)]:
            if value:
                properties.set(key, value)
        if nearby:
            root.find(f'.//{{{A}}}t').text = nearby
        if duplicate:
            clone = deepcopy(shape)
            clone.find(f'.//{{{P}}}cNvPr').set('id', '99')
            root.find(f'.//{{{P}}}spTree').append(clone)
    return changed_xml(deck(), 'ppt/slides/slide2.xml', edit)


@pytest.mark.parametrize('kwargs,status', [({}, 'missing'), ({'description': 'Synthetic caption'}, 'present'),
    ({'description': 'Synthetic caption', 'title': 'Other caption'}, 'contradictory'),
    ({'nearby': 'Caption: Synthetic caption'}, 'not_attributable'),
    ({'nearby': 'Caption: Synthetic caption\nCaption: Other caption'}, 'ambiguous')])
def test_metadata_attribution_and_no_invented_bibliography(kwargs, status):
    source = SourceDocument.read('synthetic.pptx', deck_with_caption(**kwargs), 'source')
    asset = source.metadata.assets[0]
    assert asset.fields['caption'].status == status
    assert asset.caption == ('Synthetic caption' if status == 'present' else None)
    assert asset.author is None and asset.rights is None
    assert asset.evidence_eligible is False


def test_correction_preserves_initial_values_actor_and_revision(tmp_path):
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck_with_caption(description='Synthetic caption', title='Other caption'), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    values = [{name: getattr(asset, name) for name in asset.fields} for asset in source.metadata.assets]
    values[0]['caption'] = 'Synthetic caption'
    ids = [a.id for a in source.metadata.assets]
    repository.review_assets('owner', 'source', 0, values, ids)
    values[0]['caption'] = 'Other caption'
    repository.review_assets('owner', 'source', 1, values, ids)
    restored = repository.library_resource('owner', 'source')
    first, second = restored.metadata.asset_reviews
    assert first.initial['caption'] is None and first.corrected['caption'] == 'Synthetic caption'
    assert second.initial['caption'] == 'Synthetic caption' and second.corrected['caption'] == 'Other caption'
    assert second.revision == 2 and second.actor == 'owner'
    assert SourceDocument.verify(restored, restored._original_content)


def test_identical_shapes_have_different_stable_identity():
    content = deck_with_caption(duplicate=True)
    source = SourceDocument.read('synthetic.pptx', content, 'source')
    one, two = source.metadata.assets
    assert one.sha256 == two.sha256
    assert one.id != two.id
    assert SourceDocument.read('synthetic.pptx', content, 'other-import').metadata.assets == source.metadata.assets


@pytest.mark.parametrize('surface', ['caption', 'filename', 'ocr', 'pixels', 'metadata'])
def test_sensitive_surfaces_refused_before_storage(engine_double, surface):
    from app.domain.exceptions.workflow_error import WorkflowError
    from app.tests.test_pptx_sources import attack
    from app.domain.models.source_metadata import OcrRegion
    filename = 'synthetic.pptx'
    content = deck(image=synthetic_image())
    if surface == 'caption':
        content = deck_with_caption(description='alice@example.invalid')
    elif surface == 'filename':
        filename = 'Alice Example.pptx'
    elif surface == 'ocr':
        engine_double.lines = [OcrRegion(text='alice@example.invalid', box=(0, 0, 1, 1), confidence=0.99)]
    elif surface == 'pixels':
        engine_double.faces = 1
    else:
        content = attack('metadata_email')
    with pytest.raises(WorkflowError) as error:
        SourceDocument.read(filename, content, 'source')
    assert 'alice@' not in str(error.value)


def test_pdf_original_region_preview_and_table_order(engine_double):
    from app.application.services.source_assets import asset_content
    content = table_pdf()
    with fitz.open(stream=content, filetype='pdf') as doc:
        page = doc.new_page()
        page.insert_text((40, 40), 'Publication date: 2024')
        doc.select([1, 0])
        content = doc.tobytes()
    source = SourceDocument.read('synthetic.pdf', content, 'source')
    asset, = source.metadata.assets
    assert asset.location.number == 2
    preview, media = asset_content(source, asset.id, original_region=True)
    assert preview.startswith(b'\x89PNG') and media == 'image/png'


@pytest.mark.parametrize('field,key', [('rights', 'Copyright'), ('author', 'Author'), ('caption', 'Description')])
def test_embedded_image_fields_are_attributable(engine_double, field, key):
    from io import BytesIO
    from PIL import Image, PngImagePlugin
    info = PngImagePlugin.PngInfo()
    info.add_text(key, 'synthetic public attribution')
    output = BytesIO()
    Image.new('RGB', (300, 200), 'white').save(output, format='PNG', pnginfo=info)
    source = SourceDocument.read('synthetic.pptx', deck(image=output.getvalue()), 'source')
    asset = source.metadata.assets[1]
    assert asset.fields[field].status == 'present'
    assert getattr(asset, field) == 'synthetic public attribution'
    assert asset.fields[field].origins == ['embedded_image:' + key]


def test_unattributed_or_fabricated_corrections_are_refused(tmp_path):
    from app.domain.exceptions.workflow_error import WorkflowError
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck_with_caption(nearby='Caption: Synthetic caption'), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    values = [{name: getattr(asset, name) for name in asset.fields} for asset in source.metadata.assets]
    for caption in ('Synthetic caption', 'Fabricated caption', 'alice@example.invalid'):
        values[0]['caption'] = caption
        with pytest.raises(WorkflowError):
            repository.review_assets('owner', 'source', 0, values, [a.id for a in source.metadata.assets])
    assert not repository.library_resource('owner', 'source').metadata.asset_reviews


def test_remove_keeps_library_but_permanent_deletion_purges_copies(tmp_path):
    from app.application.use_cases.manage_project_resources import RemoveProjectResourceUseCase
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    thread, state = repository.load('owner', 'demo')
    RemoveProjectResourceUseCase().execute(state, 'source')
    repository.save('owner', 'demo', thread, state)
    assert repository.library_resource('owner', 'source').metadata.assets
    assert repository.load('owner', 'demo')[1].resource_library == []
    repository.permanently_delete_source('owner', 'source', 'owner')
    assert repository.list_library_resources('owner') == []


def test_shape_order_and_slide_order_define_original_location():
    from app.tests.test_pptx_sources import changed_xml
    from app.application.services.pptx_document import P, A
    content = deck()
    content = changed_xml(content, 'ppt/slides/slide2.xml', lambda root: setattr(root.find(f'.//{{{A}}}t'), 'text', 'Publication date: 2024'))
    def reverse(root):
        slides = root.find(f'{{{P}}}sldIdLst')
        slides[:] = list(reversed(list(slides)))
    content = changed_xml(content, 'ppt/presentation.xml', reverse)
    source = SourceDocument.read('synthetic.pptx', content, 'source')
    assert source.metadata.assets[0].location.number == 1
    assert '25' in source.extracted_pages[0]['text']


def test_pdf_table_source_date_is_required_and_conflicts_fail(engine_double):
    from app.domain.exceptions.workflow_error import WorkflowError
    for text in ('No scientific date', 'Published: 2023\nPublished: 2024'):
        with fitz.open(stream=table_pdf(), filetype='pdf') as doc:
            page = doc.new_page()
            page.insert_text((40, 40), text)
            doc.select([1, 0])
            with pytest.raises(WorkflowError) as error:
                SourceDocument.read('synthetic.pdf', doc.tobytes(), 'source')
        assert error.value.code in ('SOURCE_DATE_REQUIRED', 'SOURCE_DATE_CONFLICT')


def test_051_inventory_upgrade_preserves_original_and_is_atomic(tmp_path, monkeypatch):
    import json
    from io import BytesIO
    from pptx import Presentation
    from app.domain.models.source_metadata import SourceAsset
    from app.tests.test_source_lifecycle import snapshot
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    source = SourceDocument.read('synthetic.pptx', deck(), 'source')
    repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    shape = next(s for s in Presentation(BytesIO(source._original_content)).slides[1].shapes if s.has_table)
    legacy = SourceAsset(id=f'slide-2-shape-{shape.shape_id}', kind='table', location=source.metadata.assets[0].location)
    source.metadata.assets = [legacy]
    with repository._connect() as connection:
        connection.execute('UPDATE owner_resources SET resource_json = ?', (source.model_dump_json(),))
        for (payload,) in connection.execute('SELECT metadata_json FROM project_resources').fetchall():
            payload = json.loads(payload)
            payload['metadata'] = source.metadata.model_dump(mode='json')
            connection.execute('UPDATE project_resources SET metadata_json = ?', (json.dumps(payload),))
    before = snapshot(repository)
    original = repository._insert_event
    def fail(*args):
        raise RuntimeError('Synthetic audit fault')
    monkeypatch.setattr(repository, '_insert_event', fail)
    with pytest.raises(RuntimeError):
        repository.prepare_asset_review('owner', 'source')
    assert snapshot(repository) == before
    monkeypatch.setattr(repository, '_insert_event', original)
    repository.prepare_asset_review('owner', 'source')
    restored = UserSessionRepository(repository.database_path).library_resource('owner', 'source')
    assert restored.metadata.assets[0].content_base64
    assert restored._original_content == source._original_content
    assert restored.metadata.assets[0].review_status == 'pending'


def test_derived_budget_failure_leaves_database_unchanged(tmp_path, monkeypatch):
    from app.application.services import source_assets
    from app.domain.exceptions.workflow_error import WorkflowError
    from app.tests.test_source_lifecycle import snapshot
    repository = UserSessionRepository(tmp_path / 'assets.sqlite3')
    before = snapshot(repository)
    monkeypatch.setattr(source_assets, 'MAX_DERIVED_BYTES', 1)
    with pytest.raises(WorkflowError):
        source = SourceDocument.read('synthetic.pptx', deck(), 'source')
        repository.save('owner', 'demo', 'thread', GraphState(resource_library=[source]))
    assert snapshot(repository) == before


def test_asset_query_validation_does_not_echo_sensitive_input(workspace):
    client, _ = workspace
    identifier = client.post('/resources/synthetic-owner/demo', files={'file': ('synthetic.pptx', deck())}).json()['resource_id']
    path = f'/users/synthetic-owner/resources/{identifier}'
    asset = client.get(path + '/asset-review').json()['assets'][0]
    error = client.get(path + '/assets/' + asset['id'], params={'original_region': 'alice@example.invalid'})
    assert error.status_code == 422
    assert 'alice@' not in error.text
