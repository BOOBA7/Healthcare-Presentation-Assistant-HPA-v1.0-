"""Opt-in 10.1 native rendering check using synthetic public-safe content.

Run on macOS with LibreOffice installed:
  venv/bin/python -m app.tests.evaluate_presentation_rendering

The command writes only to a temporary directory and reports a structured
unavailable-tool error instead of treating a missing renderer as a pass.
"""

import json
from pathlib import Path
from tempfile import TemporaryDirectory

import fitz
from pptx import Presentation
from pptx.util import Inches, Pt

from app.application.services.presentation_rendering import LocalPresentationRenderer
from app.domain.exceptions.workflow_error import WorkflowError


def _synthetic_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(5.8)).text_frame
    frame.text = "Synthetic public-safe rendering fixture"
    frame.paragraphs[0].font.size = Pt(30)
    paragraph = frame.add_paragraph()
    paragraph.text = "One editable PPTX source produces the reviewed PDF."
    paragraph.font.size = Pt(20)
    deck.save(path)


def evaluate() -> dict[str, object]:
    with TemporaryDirectory(prefix="hpa-render-evaluation-") as temporary:
        workspace = Path(temporary)
        pptx_path = workspace / "synthetic-rendering-fixture.pptx"
        pdf_path = workspace / "synthetic-rendering-fixture.pdf"
        _synthetic_deck(pptx_path)
        editable = Presentation(pptx_path)
        source_text = "\n".join(
            shape.text for slide in editable.slides for shape in slide.shapes if hasattr(shape, "text")
        )
        LocalPresentationRenderer().render_pdf(pptx_path, pdf_path)
        with fitz.open(pdf_path) as document:
            pdf_text = "\n".join(page.get_text() for page in document)
            page_count = document.page_count
            page_rect = document[0].rect
        expected = "Synthetic public-safe rendering fixture"
        if len(editable.slides) != 1 or expected not in source_text:
            raise AssertionError("The editable PPTX fixture was not preserved.")
        if page_count != 1 or expected not in pdf_text:
            raise AssertionError("The PDF did not preserve the synthetic slide text and order.")
        return {
            "source": "single editable synthetic PPTX",
            "renderer": "local LibreOffice headless",
            "pptx_editable": True,
            "pdf_pages": page_count,
            "pdf_landscape": page_rect.width > page_rect.height,
            "text_preserved": True,
            "limits": [
                "Manual visual comparison is still required for clipping and overlap.",
                "Font substitution and advanced PowerPoint features may differ.",
                "This fixture does not exercise production export eligibility gates.",
            ],
        }


if __name__ == "__main__":
    try:
        print(json.dumps(evaluate(), indent=2))
    except WorkflowError as error:
        print(json.dumps({"ok": False, "code": error.code, "error": error.user_message}, indent=2))
        raise SystemExit(2)
